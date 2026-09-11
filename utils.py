import os
import io
import csv
import zipfile
import openpyxl
import datetime
from PIL import Image, ImageDraw, ImageFont

LOG_FILE = "ocr_processing_log.csv"

def pdf_to_images(pdf_bytes: bytes, dpi=300, poppler_path=None) -> list:
    """
    Converts PDF bytes into a list of high-resolution PIL Images.
    Uses pdf2image by default (requires Poppler) and falls back to PyMuPDF automatically.
    """
    try:
        import pdf2image
        # If poppler_path is empty string, convert to None
        p_path = poppler_path if poppler_path and poppler_path.strip() else None
        images = pdf2image.convert_from_bytes(pdf_bytes, dpi=dpi, poppler_path=p_path)
        print("PDF converted to images using pdf2image successfully.")
        return images
    except Exception as e:
        print(f"pdf2image failed (likely due to missing Poppler): {e}. Falling back to PyMuPDF (fitz)...")
        try:
            import fitz
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            images = []
            zoom = dpi / 72  # 72 is standard PDF DPI
            mat = fitz.Matrix(zoom, zoom)
            for page in doc:
                pix = page.get_pixmap(matrix=mat)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                images.append(img)
            doc.close()
            print("PDF converted to images using PyMuPDF successfully.")
            return images
        except Exception as py_e:
            print(f"PyMuPDF fallback failed: {py_e}")
            raise RuntimeError(f"Failed to convert PDF. Ensure PyMuPDF is installed or Poppler is configured: {py_e}")

def has_border(cell) -> bool:
    """
    Checks if an Excel cell has styled borders (excluding default styleless borders).
    """
    try:
        b = cell.border
        if b is None:
            return False
        for side in [b.top, b.bottom, b.left, b.right]:
            if side and side.style is not None:
                return True
    except Exception:
        pass
    return False

def find_target_cell_to_right(sheet, row: int, col: int) -> str:
    """
    Scans cells to the right of an anchor label in Excel to find the first target cell.
    Supports merged cell ranges and skips spacer columns.
    """
    merged_ranges = sheet.merged_cells.ranges
    def get_merged_range(r, c):
        for rng in merged_ranges:
            if r >= rng.min_row and r <= rng.max_row and c >= rng.min_col and c <= rng.max_col:
                return rng
        return None
        
    anchor_range = get_merged_range(row, col)
    start_col = anchor_range.max_col + 1 if anchor_range else col + 1
    max_c = min(sheet.max_column, 50)
    
    candidates = []
    for c in range(start_col, max_c + 1):
        col_letter = openpyxl.utils.get_column_letter(c)
        
        # Skip spacer columns
        try:
            col_width = sheet.column_dimensions[col_letter].width
            if col_width is not None and col_width < 3.0:
                continue
        except Exception:
            pass
            
        cell = sheet.cell(row=row, column=c)
        rng = get_merged_range(row, c)
        
        if rng:
            top_left = sheet.cell(row=rng.min_row, column=rng.min_col)
            if top_left.value is None or str(top_left.value).strip() == '':
                return top_left.coordinate
            continue
        else:
            if cell.value is None or str(cell.value).strip() == '':
                candidates.append(cell)
                if has_border(cell):
                    return cell.coordinate
                    
    if candidates:
        return candidates[0].coordinate
    return None

def find_target_cell_below(sheet, row: int, col: int) -> str:
    """
    Scans cells below an anchor label in Excel to find the first target cell.
    Supports merged cell ranges and skips spacer rows.
    """
    merged_ranges = sheet.merged_cells.ranges
    def get_merged_range(r, c):
        for rng in merged_ranges:
            if r >= rng.min_row and r <= rng.max_row and c >= rng.min_col and c <= rng.max_col:
                return rng
        return None
        
    anchor_range = get_merged_range(row, col)
    start_row = anchor_range.max_row + 1 if anchor_range else row + 1
    max_r = min(sheet.max_row, 100)
    
    candidates = []
    for r in range(start_row, max_r + 1):
        # Skip spacer rows
        try:
            row_height = sheet.row_dimensions[r].height
            if row_height is not None and row_height < 8.0:
                continue
        except Exception:
            pass
            
        cell = sheet.cell(row=r, column=col)
        rng = get_merged_range(r, col)
        
        if rng:
            top_left = sheet.cell(row=rng.min_row, column=rng.min_col)
            if top_left.value is None or str(top_left.value).strip() == '':
                return top_left.coordinate
            continue
        else:
            if cell.value is None or str(cell.value).strip() == '':
                candidates.append(cell)
                if has_border(cell):
                    return cell.coordinate
                    
    if candidates:
        return candidates[0].coordinate
    return None

def parse_excel_template(excel_bytes: bytes) -> dict:
    """
    Parses Excel template to locate potential anchor labels and identify matching targets (Form Mode).
    """
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheets_info = {}
    for name in wb.sheetnames:
        sheet = wb[name]
        anchors = []
        max_r = min(sheet.max_row, 100)
        max_c = min(sheet.max_column, 50)
        
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                cell = sheet.cell(row=r, column=c)
                val = cell.value
                
                # Check for label cells (strings, not empty, not formulas)
                if isinstance(val, str) and len(val.strip()) > 1 and not val.strip().startswith('='):
                    val_clean = val.strip()
                    
                    right_target = find_target_cell_to_right(sheet, r, c)
                    below_target = find_target_cell_below(sheet, r, c)
                    
                    if right_target or below_target:
                        anchors.append({
                            'cell': cell.coordinate,
                            'value': val_clean,
                            'row': r,
                            'col': c,
                            'right_target': right_target,
                            'below_target': below_target
                        })
        sheets_info[name] = anchors
    return sheets_info

def read_excel_headers(excel_bytes: bytes) -> dict:
    """
    Reads the column headers from the first row of each sheet in the Excel template (Table Mode).
    """
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheets_headers = {}
    for name in wb.sheetnames:
        sheet = wb[name]
        headers = []
        max_c = min(sheet.max_column, 50)
        for col in range(1, max_c + 1):
            val = sheet.cell(row=1, column=col).value
            if val is not None:
                headers.append({
                    'col': col,
                    'name': str(val).strip(),
                    'cell_letter': openpyxl.utils.get_column_letter(col)
                })
        sheets_headers[name] = headers
    return sheets_headers

def fill_excel_form(excel_bytes: bytes, mappings: list) -> bytes:
    """
    Fills specific cell coordinates in Excel template and returns the updated file bytes (Form Mode).
    """
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    for m in mappings:
        sheet_name = m['sheet']
        cell_coord = m['cell']
        new_val = str(m['value']).strip()
        
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            
            # DataType casting
            if new_val == "":
                sheet[cell_coord] = None
            else:
                try:
                    if '.' in new_val:
                        sheet[cell_coord] = float(new_val)
                    else:
                        sheet[cell_coord] = int(new_val)
                except ValueError:
                    sheet[cell_coord] = new_val
                    
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def fill_excel_table(excel_bytes: bytes, mappings_by_pdf: dict) -> bytes:
    """
    Appends a new row of data per PDF file under matching column headers (Table Mode).
    """
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    for name in wb.sheetnames:
        sheet = wb[name]
        
        # Read column headers
        headers = {}
        for col in range(1, min(sheet.max_column, 50) + 1):
            val = sheet.cell(row=1, column=col).value
            if val is not None:
                headers[str(val).strip().lower()] = col
                
        # Find the actual last populated row
        last_row = 1
        for r in range(sheet.max_row, 0, -1):
            has_val = False
            for c in range(1, min(sheet.max_column, 50) + 1):
                if sheet.cell(row=r, column=c).value is not None:
                    has_val = True
                    break
            if has_val:
                last_row = r
                break
                
        start_row = last_row + 1
        
        # Append rows
        for pdf_name, sheet_data in mappings_by_pdf.items():
            if name in sheet_data:
                record = sheet_data[name]
                if not record:
                    continue
                    
                # Fill cells
                for header_name, val in record.items():
                    col_idx = headers.get(header_name.lower())
                    if col_idx:
                        cell = sheet.cell(row=start_row, column=col_idx)
                        val_str = str(val).strip()
                        if val_str == "":
                            cell.value = None
                        else:
                            try:
                                if '.' in val_str:
                                    cell.value = float(val_str)
                                else:
                                    cell.value = int(val_str)
                            except ValueError:
                                cell.value = val_str
                start_row += 1
                
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def log_ocr_transaction(file_name: str, page_num: int, extracted_text: str, corrected_text: str, confidence_score: float):
    """
    Logs an OCR and correction event to a local CSV database.
    """
    file_exists = os.path.exists(LOG_FILE)
    try:
        with open(LOG_FILE, 'a', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            if not file_exists:
                writer.writerow(["Timestamp", "File Name", "Page Number", "Extracted Text", "Corrected Text", "Confidence Score"])
            writer.writerow([
                datetime.datetime.now().isoformat(),
                file_name,
                page_num,
                extracted_text,
                corrected_text,
                confidence_score
            ])
    except Exception as e:
        print(f"Logging failed: {e}")

def create_export_zip(excel_filename: str, excel_bytes: bytes, original_pdfs_dict: dict, processed_images_dict: dict) -> bytes:
    """
    Creates a ZIP file containing the output Excel, original PDFs, highlighted images, and log files.
    """
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        # 1. Excel File
        zip_file.writestr(excel_filename, excel_bytes)
        
        # 2. Original PDFs
        for name, pdf_bytes in original_pdfs_dict.items():
            zip_file.writestr(f"original_pdfs/{name}", pdf_bytes)
            
        # 3. Processed Images
        for key, img in processed_images_dict.items():
            img_byte_arr = io.BytesIO()
            img.save(img_byte_arr, format='PNG')
            zip_file.writestr(f"processed_images/{key}", img_byte_arr.getvalue())
            
        # 4. Log CSV
        if os.path.exists(LOG_FILE):
            try:
                with open(LOG_FILE, 'rb') as f:
                    zip_file.writestr("ocr_processing_log.csv", f.read())
            except Exception:
                pass
                
        # 5. Learning Database
        if os.path.exists("learning_data.json"):
            try:
                with open("learning_data.json", 'rb') as f:
                    zip_file.writestr("learning_data.json", f.read())
            except Exception:
                pass
                
    return zip_buffer.getvalue()

def draw_visual_highlights(img, unified_blocks, page_mappings, draw_gray_layout=True) -> Image.Image:
    """
    Draws blue boxes on anchors, green boxes on auto-detected values, and red boxes on manual crop values.
    If draw_gray_layout is False, it skips background layout boxes to keep the original page clean.
    """
    draw_img = img.copy()
    draw = ImageDraw.Draw(draw_img, "RGBA")
    
    # Load font
    try:
        font = ImageFont.truetype("tahoma.ttf", 16)
    except IOError:
        font = ImageFont.load_default()
        
    # Draw all detected blocks in thin light gray (background layout) if requested
    if draw_gray_layout:
        for block in unified_blocks:
            x1, y1, x2, y2 = block['bbox']
            draw.rectangle([x1, y1, x2, y2], outline=(148, 163, 184, 80), fill=(148, 163, 184, 15), width=1)
            
    # Draw matched anchors in blue and values in green/red
    for m in page_mappings:
        anchor_rect = m.get('anchor_rect')
        if anchor_rect:
            ax1 = anchor_rect['x']
            ay1 = anchor_rect['y']
            ax2 = ax1 + anchor_rect['width']
            ay2 = ay1 + anchor_rect['height']
            draw.rectangle([ax1, ay1, ax2, ay2], outline=(99, 102, 241, 255), fill=(99, 102, 241, 35), width=2)
            
        value_rect = m.get('value_rect')
        if value_rect:
            vx1 = value_rect['x']
            vy1 = value_rect['y']
            vx2 = vx1 + value_rect['width']
            vy2 = vy1 + value_rect['height']
            
            # Use red color for manually defined/adjusted crops, green for auto-detected crops
            is_manual = m.get('is_manual_crop', False)
            border_color = (239, 68, 68, 255) if is_manual else (34, 197, 94, 255)
            fill_color = (239, 68, 68, 35) if is_manual else (34, 197, 94, 35)
            
            draw.rectangle([vx1, vy1, vx2, vy2], outline=border_color, fill=fill_color, width=2)
            
            # Label badge (dynamic width based on text length)
            label_text = m.get('cell') or m.get('header') or "Value"
            badge_w = len(label_text) * 9 + 10
            draw.rectangle([vx1, max(0, vy1 - 20), vx1 + badge_w, vy1], fill=border_color)
            draw.text((vx1 + 5, max(0, vy1 - 18)), str(label_text), fill=(255, 255, 255, 255), font=font)
            
    return draw_img

