import os
import re
import io
import fitz  # PyMuPDF
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image

try:
    from winocr import recognize_pil_sync as winocr_recognize
    WINOCR_AVAILABLE = True
except Exception:
    winocr_recognize = None
    WINOCR_AVAILABLE = False

def extract_quotation_data(pdf_bytes: bytes) -> dict:
    """
    Extract quotation metadata and itemized serial number list from a Quotation PDF.
    Supports Thai Heart Calibration and general calibration supplier quote formats.
    Includes OCR fallback for scanned PDFs.
    """
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    all_text = ""
    pages_text = []
    
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text = page.get_text("text")
        
        # Scanned PDF Fallback
        if len(text.strip()) < 15 and WINOCR_AVAILABLE:
            try:
                pix = page.get_pixmap(dpi=200)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_res = winocr_recognize(img)
                text = ocr_res.get("text", "")
            except Exception:
                pass
                
        pages_text.append(text)
        all_text += f"\n--- Page {page_num+1} ---\n" + text
        
    doc.close()
    all_text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', all_text)
    
    # Extract Header Metadata
    quotation_no = None
    q_match = re.search(r'(?:Quotation\s*No\.?|ใบเสนอราคาเลขที่)\s*[:/]?\s*([0-9A-Za-z\-_/]+)', all_text, re.IGNORECASE)
    if q_match:
        quotation_no = q_match.group(1).strip()
        
    quote_date = None
    d_match = re.search(r'(?:Date|วันที่)\s*[:/]?\s*([0-9]{1,2}[/\-.][0-9]{1,2}[/\-.][0-9]{2,4})', all_text, re.IGNORECASE)
    if d_match:
        quote_date = d_match.group(1).strip()
        
    customer_name = None
    c_match = re.search(r'(?:Customer|ลูกค้า|ลูกค้า\s*/\s*Customer)\s*[:/]?\s*([^\n\r]+)', all_text, re.IGNORECASE)
    if c_match:
        customer_name = c_match.group(1).strip()
        
    project_ref = None
    p_match = re.search(r'(?:Project|โครงการ|Ref\.?\s*No\.?)\s*[:/]?\s*([^\n\r]+)', all_text, re.IGNORECASE)
    if p_match:
        project_ref = p_match.group(1).strip()

    # Parse Itemized Serial Numbers
    # Standard format: S/N:56042-1915766-1170-1 (4225490) Range 0-10 bar
    items = []
    seen_serials = set()
    
    # Regex for structured S/N lines
    sn_pattern = re.compile(
        r'S/N\s*[:\s]\s*([A-Za-z0-9\-_.]+)(?:\s*\(([^)]+)\))?(?:\s*Range\s*([^\n\r,]+))?',
        re.IGNORECASE
    )
    
    current_category = ""
    lines = all_text.splitlines()
    
    for line in lines:
        line_clean = line.strip()
        if not line_clean:
            continue
            
        # Check category headers
        cat_match = re.match(r'^(?:\d+\s*/)?\s*([A-Za-z\s]+)(?:/\s*\d+\s*unit)?', line_clean)
        if cat_match and not line_clean.startswith("S/N") and "Calibration" not in line_clean and len(line_clean) < 50:
            candidate_cat = cat_match.group(1).strip()
            if candidate_cat.lower() not in ["date", "attn", "tel", "fax", "no", "page", "total", "description"]:
                current_category = candidate_cat
                
        # Look for S/N matches
        match = sn_pattern.search(line_clean)
        if match:
            serial_no = match.group(1).strip()
            comp_no = match.group(2).strip() if match.group(2) else ""
            range_val = match.group(3).strip() if match.group(3) else ""
            
            if not range_val and "Range" in line_clean:
                r_part = line_clean.split("Range", 1)[1].strip()
                range_val = r_part
                
            norm_key = re.sub(r'[^A-Za-z0-9]', '', serial_no).upper()
            if norm_key and norm_key not in seen_serials:
                seen_serials.add(norm_key)
                items.append({
                    "Quotation S/N": serial_no,
                    "Normalized S/N": norm_key,
                    "Quotation Component No": comp_no,
                    "Quotation Description": current_category,
                    "Quotation Range / Specs": range_val,
                    "Raw Line": line_clean
                })
                
    # Fallback scan for standard machine serial pattern: 56XXX-XXXXXXX-XXXX-X or 56XXX-XXXXXXX-XX-X
    if not items:
        standalone_sn_pattern = re.compile(r'\b(5\d{4}-\d+-\d+(?:-\d+)?)\b')
        for line in lines:
            line_clean = line.strip()
            for m in standalone_sn_pattern.finditer(line_clean):
                serial_no = m.group(1).strip()
                norm_key = re.sub(r'[^A-Za-z0-9]', '', serial_no).upper()
                if norm_key not in seen_serials:
                    seen_serials.add(norm_key)
                    c_m = re.search(r'\((\d{7,8})\)', line_clean)
                    comp_no = c_m.group(1) if c_m else ""
                    items.append({
                        "Quotation S/N": serial_no,
                        "Normalized S/N": norm_key,
                        "Quotation Component No": comp_no,
                        "Quotation Description": "",
                        "Quotation Range / Specs": "",
                        "Raw Line": line_clean
                    })
                    
    df_items = pd.DataFrame(items) if items else pd.DataFrame(columns=[
        "Quotation S/N", "Normalized S/N", "Quotation Component No", "Quotation Description", "Quotation Range / Specs", "Raw Line"
    ])
    
    return {
        "quotation_no": quotation_no or "N/A",
        "date": quote_date or "N/A",
        "customer": customer_name or "N/A",
        "project_ref": project_ref or "N/A",
        "total_items": len(df_items),
        "items_df": df_items,
        "raw_text": all_text
    }

def extract_calibration_list_data(source, filename: str = "") -> pd.DataFrame:
    """
    Extract itemized calibration records (Serial No, Component No, Description, Size/dimensions)
    from a pandas DataFrame, Excel bytes/filepath, or Calibration PDF bytes.
    Supports processed calibration sheets, raw BOM Excel sheets, and converted PDFs.
    """
    df = None
    
    if isinstance(source, pd.DataFrame):
        df = source.copy()
    elif isinstance(source, bytes):
        if source.startswith(b'%PDF'):
            doc = fitz.open(stream=source, filetype="pdf")
            records = []
            prev_cols = {}
            
            # Step 1: Structured table extraction across pages
            for p in range(len(doc)):
                page = doc[p]
                try:
                    tables = page.find_tables()
                    for tab in tables:
                        df_tab = tab.extract()
                        if not df_tab or len(df_tab) < 1:
                            continue
                        
                        first_row_str = " ".join([str(x).lower() for x in df_tab[0] if x])
                        is_header = any(k in first_row_str for k in ["component", "serial", "s/n", "object", "description"])
                        
                        if is_header:
                            header = [str(h).replace("\n", " ").strip().lower() for h in df_tab[0]]
                            prev_cols = {"sn": -1, "comp": -1, "desc": -1, "size": -1}
                            for idx, h in enumerate(header):
                                if "serial" in h or "s/n" in h:
                                    prev_cols["sn"] = idx
                                elif "component" in h and "number" in h:
                                    prev_cols["comp"] = idx
                                elif "description" in h:
                                    prev_cols["desc"] = idx
                                elif "size" in h or "dimension" in h:
                                    prev_cols["size"] = idx
                            data_rows = df_tab[1:]
                        else:
                            data_rows = df_tab
                            
                        sn_idx = prev_cols.get("sn", -1)
                        comp_idx = prev_cols.get("comp", -1)
                        desc_idx = prev_cols.get("desc", -1)
                        size_idx = prev_cols.get("size", -1)
                        
                        for row in data_rows:
                            sn_val = ""
                            if sn_idx != -1 and sn_idx < len(row) and row[sn_idx]:
                                sn_val = str(row[sn_idx]).strip()
                            if not sn_val or not re.search(r'\b5\d{4}-', sn_val):
                                for c in row:
                                    m = re.search(r'\b(5\d{4}(?:-[A-Za-z0-9]+){2,4})\b', str(c or ""))
                                    if m:
                                        sn_val = m.group(1)
                                        break
                                        
                            if sn_val and ("-" in sn_val or re.search(r'\d', sn_val)) and len(sn_val) >= 5:
                                comp = str(row[comp_idx]).strip().replace("\n", " ") if comp_idx != -1 and comp_idx < len(row) and row[comp_idx] else ""
                                desc = str(row[desc_idx]).strip().replace("\n", " ") if desc_idx != -1 and desc_idx < len(row) and row[desc_idx] else ""
                                size = str(row[size_idx]).strip().replace("\n", " ") if size_idx != -1 and size_idx < len(row) and row[size_idx] else ""
                                records.append({
                                    "Serial No.": sn_val,
                                    "Component number": comp,
                                    "Object description": desc,
                                    "Size/dimensions": size
                                })
                except Exception:
                    pass
            
            # Step 2: Fallback line extraction if no table records found
            if not records:
                pdf_lines = []
                for p in range(len(doc)):
                    pdf_lines.extend(doc.load_page(p).get_text("text").splitlines())
                    
                sn_regex = re.compile(r'\b(5\d{4}(?:-[A-Za-z0-9]+){2,4})\b')
                for l in pdf_lines:
                    m = sn_regex.search(l)
                    if m:
                        records.append({
                            "Serial No.": m.group(1).strip(),
                            "Component number": "",
                            "Object description": "",
                            "Size/dimensions": ""
                        })
            doc.close()
            df = pd.DataFrame(records)
        else:
            try:
                df = pd.read_excel(io.BytesIO(source))
            except Exception:
                xl = pd.ExcelFile(io.BytesIO(source))
                for sheet in xl.sheet_names:
                    temp_df = xl.parse(sheet)
                    if any("serial" in str(c).lower() for c in temp_df.columns):
                        df = temp_df
                        break
    elif isinstance(source, str) and os.path.exists(source):
        filename = filename or os.path.basename(source)
        df = pd.read_excel(source)
        
    if df is None or df.empty:
        return pd.DataFrame(columns=["List S/N", "Normalized S/N", "List Component No", "List Description", "List Size / Dimensions"])
        
    # Ensure all column names are unique string identifiers to avoid DataFrame indexing issues
    cols = []
    counts = {}
    for c in df.columns:
        name = str(c).strip()
        if name in counts:
            counts[name] += 1
            cols.append(f"{name}_{counts[name]}")
        else:
            counts[name] = 0
            cols.append(name)
    df.columns = cols

    # Case A: Check if this is a raw BOM Excel sheet with CALIBRATION rows where serials need to be compiled
    has_calib_col = any("internal processes" in str(c).lower() for c in df.columns)
    has_serial_col = any("serial" in str(c).lower() or "s/n" in str(c).lower() for c in df.columns)
    
    if has_calib_col and not has_serial_col:
        calib_col = next(c for c in df.columns if "internal processes" in str(c).lower())
        proj_code = "5XXXX"
        if filename:
            m_proj = re.search(r'(\d{5})', filename)
            if m_proj:
                proj_code = m_proj.group(1)
        if proj_code == "5XXXX":
            for c in df.columns[:5]:
                sample_s = df[c].dropna().astype(str)
                m_cell = sample_s[sample_s.str.contains(r'\b5\d{4}\b', regex=True)]
                if len(m_cell) > 0:
                    proj_code = re.search(r'\b(5\d{4})\b', m_cell.iloc[0]).group(1)
                    break
                    
        level_col = df.columns[0]
        item_col = df.columns[1]
        comp_col = df.columns[2]
        desc_col = df.columns[3] if len(df.columns) > 3 else comp_col
        size_col = df.columns[4] if len(df.columns) > 4 else comp_col
        qty_col = df.columns[6] if len(df.columns) > 6 else None
        
        last_level2_comp = ""
        pc_counter = {}
        raw_calib_rows = []
        
        for _, row in df.iterrows():
            lvl = str(row[level_col]).strip()
            if lvl == "2":
                last_level2_comp = str(row[comp_col]).strip()
                
            p_val = str(row.get(calib_col, "") or "").strip().upper()
            if p_val == "CALIBRATION":
                z_val = str(row[item_col]).strip()
                y_val = last_level2_comp
                comp_num = str(row[comp_col] or "").strip()
                desc = str(row[desc_col] or "").strip()
                size = str(row[size_col] or "").strip()
                qty = 1
                if qty_col:
                    try:
                        qty = int(float(row[qty_col]))
                    except Exception:
                        pass
                for _ in range(qty):
                    key = (y_val, z_val)
                    pc_counter[key] = pc_counter.get(key, 0) + 1
                    sn = f"{proj_code}-{y_val}-{z_val}-{pc_counter[key]}"
                    raw_calib_rows.append({
                        "List S/N": sn,
                        "Normalized S/N": re.sub(r'[^A-Za-z0-9]', '', sn).upper(),
                        "List Component No": comp_num,
                        "List Description": desc,
                        "List Size / Dimensions": size
                    })
        if raw_calib_rows:
            return pd.DataFrame(raw_calib_rows)

    # Case B: Standard / Processed Calibration Sheet
    col_map = {}
    assigned_targets = set()
    for col in df.columns:
        c_str = str(col).strip().lower()
        if ("serial" in c_str or "s/n" in c_str or c_str == "sn") and "Serial No." not in assigned_targets:
            col_map[col] = "Serial No."
            assigned_targets.add("Serial No.")
        elif ("component number" in c_str or "comp. no" in c_str or "comp no" in c_str or c_str == "component no" or c_str in ["comp.", "component"]) and "Component number" not in assigned_targets:
            col_map[col] = "Component number"
            assigned_targets.add("Component number")
        elif ("object description" in c_str or "description" in c_str or "item name" in c_str) and "Object description" not in assigned_targets:
            col_map[col] = "Object description"
            assigned_targets.add("Object description")
        elif ("size" in c_str or "dimension" in c_str) and "Size/dimensions" not in assigned_targets:
            col_map[col] = "Size/dimensions"
            assigned_targets.add("Size/dimensions")
        elif ("no." in c_str or c_str == "no") and "Item No." not in assigned_targets:
            col_map[col] = "Item No."
            assigned_targets.add("Item No.")
            
    df_renamed = df.rename(columns=col_map)
    
    if "Serial No." not in df_renamed.columns:
        for c in df_renamed.columns:
            series = df_renamed[c]
            if isinstance(series, pd.DataFrame):
                series = series.iloc[:, 0]
            sample_vals = series.dropna().astype(str).tolist()
            if any(re.search(r'5\d{4}-\d+-\d+', v) for v in sample_vals[:20]):
                df_renamed = df_renamed.rename(columns={c: "Serial No."})
                break
                
    if "Serial No." not in df_renamed.columns:
        return pd.DataFrame(columns=["List S/N", "Normalized S/N", "List Component No", "List Description", "List Size / Dimensions"])
        
    valid_rows = []
    for idx, row in df_renamed.iterrows():
        sn_val = str(row.get("Serial No.", "") or "").strip()
        if not sn_val or sn_val.lower() in ["nan", "none", "<na>", "serial no.", "serial no", "s/n"]:
            continue
            
        if len(sn_val) >= 5 and ("-" in sn_val or re.search(r'\d', sn_val)):
            norm_sn = re.sub(r'[^A-Za-z0-9]', '', sn_val).upper()
            comp_num = str(row.get("Component number", "") or "").replace("nan", "").replace("None", "").strip()
            desc = str(row.get("Object description", "") or "").replace("nan", "").replace("None", "").strip()
            size_dim = str(row.get("Size/dimensions", "") or "").replace("nan", "").replace("None", "").strip()
            
            valid_rows.append({
                "List S/N": sn_val,
                "Normalized S/N": norm_sn,
                "List Component No": comp_num,
                "List Description": desc,
                "List Size / Dimensions": size_dim
            })
            
    df_result = pd.DataFrame(valid_rows) if valid_rows else pd.DataFrame(columns=[
        "List S/N", "Normalized S/N", "List Component No", "List Description", "List Size / Dimensions"
    ])
    
    if not df_result.empty:
        df_result = df_result.drop_duplicates(subset=["Normalized S/N"]).reset_index(drop=True)
        
    return df_result

def compare_quotation_and_calibration(quotation_data: dict, calibration_df: pd.DataFrame) -> dict:
    """
    Compare Quotation PDF items against Calibration Certificate list items.
    Generates full matching breakdown, status flags, and discrepancy analysis.
    """
    df_quote = quotation_data["items_df"]
    df_calib = calibration_df
    
    quote_map = {}
    if not df_quote.empty:
        for _, r in df_quote.iterrows():
            quote_map[r["Normalized S/N"]] = r.to_dict()
            
    calib_map = {}
    calib_order_keys = []
    if not df_calib.empty:
        for _, r in df_calib.iterrows():
            k = str(r["Normalized S/N"]).strip().upper()
            if k not in calib_map:
                calib_order_keys.append(k)
            calib_map[k] = r.to_dict()
            
    comparison_rows = []
    matched_count = 0
    missing_in_quote_count = 0
    extra_in_quote_count = 0
    discrepancy_count = 0
    
    seen_keys = set()
    item_idx = 1
    
    # 1. Primary ordering: Strictly follow the inserted "2. Calibration Certificate List"
    for key in calib_order_keys:
        seen_keys.add(key)
        c_item = calib_map.get(key, {})
        in_quote = key in quote_map
        q_item = quote_map.get(key, {})
        
        serial_display = c_item.get("List S/N") or q_item.get("Quotation S/N") or key
        
        if in_quote:
            q_comp = str(q_item.get("Quotation Component No", "")).strip()
            c_comp = str(c_item.get("List Component No", "")).strip()
            if q_comp and c_comp and q_comp != c_comp:
                status = "⚠️ Discrepancy"
                notes = f"Component No Mismatch: Quote has '{q_comp}', List has '{c_comp}'"
                discrepancy_count += 1
            else:
                matched_count += 1
                status = "✅ Matched"
                notes = "Present in both Quotation and Calibration List."
        else:
            missing_in_quote_count += 1
            status = "⚠️ Missing in Quotation"
            notes = "Required by Calibration List, but NOT quoted by vendor."
            
        comparison_rows.append({
            "No.": item_idx,
            "Status": status,
            "Serial No.": serial_display,
            "List Comp No": c_item.get("List Component No", "-"),
            "Quotation Comp No": q_item.get("Quotation Component No", "-"),
            "List Desc": c_item.get("List Description", "-"),
            "Quotation Desc": q_item.get("Quotation Description", "-"),
            "List Dimensions": c_item.get("List Size / Dimensions", "-"),
            "Quotation Range / Specs": q_item.get("Quotation Range / Specs", "-"),
            "Notes": notes,
            "_norm_key": key
        })
        item_idx += 1
        
    # 2. Append any extra items in Quotation that were not in Calibration List
    for key, q_item in quote_map.items():
        if key not in seen_keys:
            extra_in_quote_count += 1
            status = "⚠️ Extra in Quotation"
            notes = "Quoted by vendor, but not in Calibration List."
            comparison_rows.append({
                "No.": item_idx,
                "Status": status,
                "Serial No.": q_item.get("Quotation S/N", key),
                "List Comp No": "-",
                "Quotation Comp No": q_item.get("Quotation Component No", "-"),
                "List Desc": "-",
                "Quotation Desc": q_item.get("Quotation Description", "-"),
                "List Dimensions": "-",
                "Quotation Range / Specs": q_item.get("Quotation Range / Specs", "-"),
                "Notes": notes,
                "_norm_key": key
            })
            item_idx += 1
        
    df_comparison = pd.DataFrame(comparison_rows) if comparison_rows else pd.DataFrame()
    
    total_calib = len(df_calib)
    total_quote = len(df_quote)
    
    match_rate = (matched_count / total_calib * 100.0) if total_calib > 0 else 0.0
    
    return {
        "quotation_meta": quotation_data,
        "total_calibration_items": total_calib,
        "total_quotation_items": total_quote,
        "matched_count": matched_count,
        "missing_in_quotation_count": missing_in_quote_count,
        "extra_in_quote_count": extra_in_quote_count,
        "discrepancy_count": discrepancy_count,
        "match_rate_pct": round(match_rate, 1),
        "comparison_df": df_comparison
    }

def generate_verification_report_excel(comp_result: dict) -> bytes:
    """
    Generate a professional styled Excel verification report with Summary KPIs and detailed comparison sheets.
    Sanitizes all cell strings using openpyxl's ILLEGAL_CHARACTERS_RE to prevent IllegalCharacterError.
    """
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    def clean_excel_val(v):
        if v is None:
            return ""
        return ILLEGAL_CHARACTERS_RE.sub("", str(v)).strip()

    wb = openpyxl.Workbook()
    
    # Sheet 1: Summary Dashboard
    ws_sum = wb.active
    ws_sum.title = "Verification Summary"
    ws_sum.views.sheetView[0].showGridLines = True
    
    font_title = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    fill_header = PatternFill("solid", fgColor="0078D4")
    font_cell = Font(name="Arial", size=10)
    font_bold = Font(name="Arial", size=10, bold=True)
    
    fill_match = PatternFill("solid", fgColor="D4EDDA")
    fill_miss = PatternFill("solid", fgColor="F8D7DA")
    fill_extra = PatternFill("solid", fgColor="FFF3CD")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    ws_sum.merge_cells("A1:F1")
    title_cell = ws_sum["A1"]
    title_cell.value = clean_excel_val("📜 CALIBRATION QUOTATION VERIFICATION REPORT")
    title_cell.font = font_title
    title_cell.fill = fill_header
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 35
    
    q_meta = comp_result.get("quotation_meta", {})
    ws_sum["A3"] = "Quotation No:"
    ws_sum["B3"] = clean_excel_val(q_meta.get("quotation_no", "N/A"))
    ws_sum["A4"] = "Quotation Date:"
    ws_sum["B4"] = clean_excel_val(q_meta.get("date", "N/A"))
    ws_sum["A5"] = "Customer / Ref:"
    ws_sum["B5"] = clean_excel_val(q_meta.get("customer", "N/A"))
    
    for r in range(3, 6):
        ws_sum[f"A{r}"].font = font_bold
        ws_sum[f"B{r}"].font = font_cell
        
    ws_sum["D3"] = "📊 Verification Metrics"
    ws_sum["D3"].font = font_bold
    ws_sum.merge_cells("D3:F3")
    ws_sum["D3"].fill = PatternFill("solid", fgColor="F2F2F2")
    
    kpis = [
        ("Total Items in Calibration List", comp_result.get("total_calibration_items", 0)),
        ("Total Items in Quotation PDF", comp_result.get("total_quotation_items", 0)),
        ("✅ Matched Serial Numbers", comp_result.get("matched_count", 0)),
        ("❌ Missing in Quotation", comp_result.get("missing_in_quotation_count", 0)),
        ("⚠️ Extra in Quotation", comp_result.get("extra_in_quotation_count", 0)),
        ("🎯 Overall Match Rate", f"{comp_result.get('match_rate_pct', 0)} %")
    ]
    
    for i, (label, val) in enumerate(kpis, start=4):
        ws_sum[f"D{i}"] = clean_excel_val(label)
        ws_sum[f"E{i}"] = clean_excel_val(val)
        ws_sum[f"D{i}"].font = font_cell
        ws_sum[f"E{i}"].font = font_bold
        ws_sum[f"D{i}"].border = thin_border
        ws_sum[f"E{i}"].border = thin_border
        
    # Sheet 2: Itemized Comparison
    ws_items = wb.create_sheet(title="Itemized Comparison")
    ws_items.views.sheetView[0].showGridLines = True
    
    headers = [
        "No.", "Status", "Serial No.", "List Component No", "Quote Component No", 
        "List Description", "Quote Description", "List Dimensions", "Quote Range / Specs", "Verification Notes"
    ]
    
    ws_items.append([clean_excel_val(h) for h in headers])
    ws_items.row_dimensions[1].height = 26
    
    for c_idx in range(1, len(headers) + 1):
        cell = ws_items.cell(row=1, column=c_idx)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    df_comp = comp_result.get("comparison_df", pd.DataFrame())
    if not df_comp.empty:
        for row_idx, r in enumerate(df_comp.iterrows(), start=2):
            data = r[1]
            status_val = str(data.get("Status", ""))
            
            row_data = [
                data.get("No.", row_idx - 1),
                status_val,
                data.get("Serial No.", ""),
                data.get("List Comp No", ""),
                data.get("Quotation Comp No", ""),
                data.get("List Desc", ""),
                data.get("Quotation Desc", ""),
                data.get("List Dimensions", ""),
                data.get("Quotation Range / Specs", ""),
                data.get("Notes", "")
            ]
            clean_row = [clean_excel_val(x) for x in row_data]
            ws_items.append(clean_row)
            
            ws_items.row_dimensions[row_idx].height = 20
            # Green if matched, Yellow if not matched
            if "Matched" in status_val and "Discrepancy" not in status_val and "Missing" not in status_val and "Extra" not in status_val:
                row_fill = fill_match # Light green
            else:
                row_fill = fill_extra # Light yellow
                
            for col_idx in range(1, len(clean_row) + 1):
                c = ws_items.cell(row=row_idx, column=col_idx)
                c.font = font_cell
                c.border = thin_border
                if row_fill:
                    c.fill = row_fill
                    
    for ws in [ws_sum, ws_items]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
            
    out_io = io.BytesIO()
    wb.save(out_io)
    return out_io.getvalue()
