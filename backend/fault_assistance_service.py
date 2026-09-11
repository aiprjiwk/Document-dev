import os
import re
import io
import zipfile
import html
import openpyxl
import pandas as pd
import logging

logger = logging.getLogger("FaultAssistanceService")

def update_project_in_index_html(content: str, project_name: str) -> str:
    content = re.sub(r'<title>.*?</title>', f'<title>{project_name}</title>', content, flags=re.IGNORECASE | re.DOTALL)
    pattern = r'(<td\s+class=["\']project["\']>\s*<a\s+href=["\']index\.html["\']>).*?(</a>)'
    content = re.sub(pattern, rf'\g<1>{project_name}\g<2>', content, flags=re.IGNORECASE | re.DOTALL)
    return content

def update_project_in_detail_html(content: str, project_name: str) -> str:
    title_match = re.search(r'<title>(.*?)</title>', content, flags=re.IGNORECASE | re.DOTALL)
    if title_match:
        old_title = title_match.group(1)
        if " - " in old_title:
            msg_part = old_title.split(" - ", 1)[1]
        else:
            msg_part = old_title
        new_title = f"{project_name} - {msg_part}"
        content = re.sub(r'<title>.*?</title>', f'<title>{new_title}</title>', content, flags=re.IGNORECASE | re.DOTALL)
    
    pattern = r'(<td\s+class=["\']project["\']>\s*<a\s+href=["\']index\.html["\']>).*?(</a>)'
    content = re.sub(pattern, rf'\g<1>{project_name}\g<2>', content, flags=re.IGNORECASE | re.DOTALL)
    return content

def process_fault_assistance_files(excel_bytes: bytes, customer_name: str, src_folder: str = "Fault assistance") -> tuple:
    try:
        # Extract project number from customer name (e.g. "56061" from "56061_Henkel...")
        project_number = customer_name.strip()
        match = re.match(r'^([A-Za-z0-9]+)', customer_name.strip())
        if match:
            project_number = match.group(1)
            
        # 1. Load Excel Worksheet
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb['stoerkenner'] if 'stoerkenner' in wb.sheetnames else wb.active
        
        # Check for Header Row (Check if A1 has "variable")
        first_cell_val = str(ws.cell(row=1, column=1).value or "").strip().lower()
        has_header = ("variable" in first_cell_val)
        start_row = 2 if has_header else 1
        
        # Write Project inputs to Excel (E2/F2 or E1/F1 based on header presence)
        if has_header:
            ws.cell(row=1, column=5, value="Project number")
            ws.cell(row=1, column=6, value="Custumer")
            ws.cell(row=2, column=5, value=project_number)  # Column E (Project number)
            ws.cell(row=2, column=6, value=customer_name)   # Column F (Custumer)
        else:
            ws.cell(row=1, column=5, value=project_number)  # Column E (Project number)
            ws.cell(row=1, column=6, value=customer_name)   # Column F (Custumer)
            
        # Determine last row based on non-empty B and C values
        last_row = ws.max_row
        while last_row >= start_row:
            val_B = ws.cell(row=last_row, column=2).value
            val_C = ws.cell(row=last_row, column=3).value
            if val_B is not None or val_C is not None:
                break
            last_row -= 1
            
        if last_row < start_row:
            return False, "No data found in column B/C from data row downward.", None, None, None, project_number, None

        # 2. Setup Zip File In Memory
        zip_buffer = io.BytesIO()
        ok_count = 0
        fail_count = 0
        status_rows = []
        fault_links = []
        xml_entries = []
        
        # 3. Copy Base Assets
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            def copy_folder_to_zip(folder_name):
                src_path = os.path.join(src_folder, folder_name)
                if os.path.exists(src_path) and os.path.isdir(src_path):
                    for root, dirs, files in os.walk(src_path):
                        for file in files:
                            full_file_path = os.path.join(root, file)
                            rel_path = os.path.relpath(full_file_path, src_path)
                            zip_path = os.path.join(folder_name, rel_path)
                            zip_file.write(full_file_path, zip_path)
            
            for folder in ["CSS", "layout", "Scripts", "Images", "safety"]:
                copy_folder_to_zip(folder)
                
            css_path = os.path.join(src_folder, "layout.css")
            if os.path.exists(css_path):
                zip_file.write(css_path, "layout.css")
                
            index_path = os.path.join(src_folder, "index.html")
            if os.path.exists(index_path):
                try:
                    with open(index_path, "r", encoding="utf-8", errors="ignore") as f:
                        index_content = f.read()
                    updated_index = update_project_in_index_html(index_content, customer_name)
                    zip_file.writestr("index.html", updated_index.encode("utf-8"))
                except Exception as index_err:
                    logger.warning(f"Error copying index.html: {index_err}")
            
            # 4. Copy and Process Row-by-Row files (Start loop from start_row)
            for r in range(start_row, last_row + 1):
                col_A_val = str(ws.cell(row=r, column=1).value or "").strip()
                col_B_val = str(ws.cell(row=r, column=2).value or "").strip()
                col_C_val = str(ws.cell(row=r, column=3).value or "").strip()
                
                # Check for empty rows to break
                if col_B_val == "" and col_C_val == "":
                    break
                
                # Add to XML entries list if Column A is not empty (Part 3)
                if col_A_val != "":
                    xml_entries.append((col_A_val, col_B_val, col_C_val))
                    
                if col_B_val == "":
                    continue
                
                file_name = f"{col_B_val}_{col_C_val}.html"
                src_file_path = os.path.join(src_folder, file_name)
                
                # Store fault link for index generation (Part 2)
                fault_links.append((file_name, f"{col_B_val}_{col_C_val}"))
                
                if os.path.exists(src_file_path):
                    try:
                        with open(src_file_path, "r", encoding="utf-8", errors="ignore") as f:
                            html_content = f.read()
                        updated_html = update_project_in_detail_html(html_content, customer_name)
                        zip_file.writestr(file_name, updated_html.encode("utf-8"))
                        
                        ok_count += 1
                        status_rows.append({"Variable": col_A_val, "Fehlern": col_B_val, "Alarmtext_Englisch": col_C_val, "Status": "OK"})
                    except Exception as copy_err:
                        fail_count += 1
                        status_rows.append({"Variable": col_A_val, "Fehlern": col_B_val, "Alarmtext_Englisch": col_C_val, "Status": f"FAILED: {str(copy_err)}"})
                else:
                    fail_count += 1
                    status_rows.append({"Variable": col_A_val, "Fehlern": col_B_val, "Alarmtext_Englisch": col_C_val, "Status": "NOT FOUND"})
            
            # 5. Generate dynamically "All fault messages.html" (Part 2 logic)
            cls_gen = "Gen55b573653f38405eadc44f09c4415b9a"
            esc_cust_name = html.escape(customer_name)
            
            html_parts = [
                '<!DOCTYPE html PUBLIC "-//W3C//DTD XHTML 1.0 Transitional//EN" "http://www.w3.org/TR/xhtml1/DTD/xhtml1-transitional.dtd">\n',
                '<html xmlns="http://www.w3.org/1999/xhtml">\n',
                '<head xmlns:forms="http://www.schema.de/2010/ST4/Layout/MarkupLanguage/Forms">\n',
                f'<title>{esc_cust_name} - All fault messages</title>\n',
                '<meta http-equiv="Content-Script-Type" content="text/javascript" />\n',
                '<meta http-equiv="Content-Style-Type" content="text/css" />\n',
                '<meta http-equiv="Content-Type" content="text/html; charset=UTF-8" />\n',
                '<meta name="generator" content="SCHEMA ST4, Bootstrap 2019 v3" />\n',
                '<script type="text/javascript" src="Scripts/Common.Core.js"></script>\n',
                '<script type="text/javascript" src="Scripts/Common.Language.js"></script>\n',
                '<script type="text/javascript" src="Scripts/Common.SearchEngine.js"></script>\n',
                '<script type="text/javascript" src="Scripts/Common.Tree.js"></script>\n',
                '<script type="text/javascript" src="Scripts/main.js"></script>\n',
                '<script type="text/javascript" src="Scripts/Production.js"></script>\n',
                '<link rel="stylesheet" type="text/css" href="layout.css" />\n',
                '<link rel="stylesheet" type="text/css" href="CSS/Common.Reset.css" />\n',
                '<link rel="stylesheet" type="text/css" href="CSS/Core.Production.css" />\n',
                '<link rel="stylesheet" type="text/css" href="CSS/HTML.css" />\n',
                '<link rel="stylesheet" type="text/css" href="CSS/main.css" />\n',
                '</head><body xmlns:forms="http://www.schema.de/2010/ST4/Layout/MarkupLanguage/Forms">\n',
                '<div class="page"><div class="company"></div>\n',
                '<table class="metatable"><tbody><tr><td class="project">\n',
                f'<a href="index.html">{esc_cust_name}</a></td><td>\n',
                '</div></td><td class="metatable togglemenu"><a href="index.html"></a></td></tr></tbody></table>\n',
                '<table class="contentLayout"><tbody><tr><td class="contentLayoutOne">\n',
                '<div class="navigation tripletNavigation"><div class="tripletBody">\n',
                '<div class="tripletMainMenu tripletMainMenuInfoType goMainMenu"><a href="index.html">Home page</a></div>\n',
                '<div class="tripletParent tripletParentInfoType goUp"><div class="tripletInfo tripletInfoParent">\n',
                '<div class="tripletImg"><img src="Images/MenuUp.png" /></div>\n',
                '<div class="tripletText level-up"><a href="index.html">One level up</a></div>\n',
                '</div></div></div><div class="tripletHead tripletHeadInfoType"></div>\n',
                '</div></div></div></div></td><td class="contentLayoutTwo"><img src="Images/Empty.gif" /></td>\n',
                '<td class="contentLayoutThree"><div class="content"><div class="anchor_clear"></div>\n',
                '<div class="anchor_margin"><div class="HeadingMargin"></div></div>\n',
                '<div class="anchor_text"><div class="Heading">All fault messages<span class="NotReleased"></span></div></div>\n',
                '<div class="metalist"></div><div class="submenu"><div class="anchor_text">\n'
            ]
            
            for file_name, link_text in fault_links:
                esc_file_name = html.escape(file_name, quote=True)
                esc_link_text = html.escape(link_text)
                html_parts.append(
                    f'<div class="anchor_text"><dl class="{cls_gen}">\n'
                    f'<dt class="{cls_gen}" style="display:none;"></dt>\n'
                    f'<dd class="{cls_gen}"><a href="{esc_file_name}">{esc_link_text}</a></dd></dl></div>\n'
                )
                
            html_parts.append('</div></div></td></tr></tbody></table></div></body></html>')
            all_fault_html = "".join(html_parts)
            zip_file.writestr("All fault messages.html", all_fault_html.encode("utf-8"))
            
            # 6. Generate dynamically "stoerkenner.xml" (Part 3 logic)
            xml_parts = [
                '<?xml version="1.0" encoding="utf-8"?>\n',
                '<MISConfiguration xmlns="http://www.w3.org/1999/xhtml">\n',
                f'\t<DocRoot>C:\\MIS-Daten\\{project_number}</DocRoot>\n',
                '\t<DocMain>index.html</DocMain>\n',
                '\t<Docs>\n'
            ]
            
            for yyyy, zzzz, xxxx in xml_entries:
                esc_yyyy = html.escape(yyyy, quote=True)
                esc_zzzz = html.escape(zzzz)
                esc_xxxx = html.escape(xxxx)
                xml_parts.append(
                    f'\t\t<DocPath MISTag="{esc_yyyy}">{esc_zzzz}_{esc_xxxx}.html</DocPath>\n'
                )
                
            xml_parts.append('\t</Docs>\n')
            xml_parts.append('</MISConfiguration>')
            
            stoerkenner_xml = "".join(xml_parts)
            zip_file.writestr("stoerkenner.xml", stoerkenner_xml.encode("utf-8"))
            
        # 7. Generate Final.zip with nested directory structure (Part 4)
        final_zip_buffer = io.BytesIO()
        with zipfile.ZipFile(final_zip_buffer, "w", zipfile.ZIP_DEFLATED) as final_zip:
            # Copy all files from primary ZIP into WithNavi and WithoutNavi directories
            zip_buffer.seek(0)
            with zipfile.ZipFile(zip_buffer, "r") as temp_zip:
                for item in temp_zip.infolist():
                    data = temp_zip.read(item.filename)
                    
                    # Write to WithNavi
                    with_navi_path = f"MIS-Daten/{project_number}/WithNavi/{item.filename}"
                    final_zip.writestr(with_navi_path, data)
                    
                    # Write to WithoutNavi
                    without_navi_path = f"MIS-Daten/{project_number}/WithoutNavi/{item.filename}"
                    final_zip.writestr(without_navi_path, data)
                    
            # Write stoerkenner.xml to MIS-Daten/{project_number}/stoerkenner.xml
            final_zip.writestr(f"MIS-Daten/{project_number}/stoerkenner.xml", stoerkenner_xml.encode("utf-8"))
        
        # Save Excel Status
        excel_out = io.BytesIO()
        wb.save(excel_out)
        
        summary_msg = f"Completed! Copied {ok_count} files, Failed/Not Found {fail_count} files, and generated All fault messages.html & stoerkenner.xml."
        status_df = pd.DataFrame(status_rows)
        
        return True, summary_msg, status_df, zip_buffer.getvalue(), excel_out.getvalue(), project_number, final_zip_buffer.getvalue()
    except Exception as e:
        return False, f"Process Error: {str(e)}", None, None, None, project_number, None
