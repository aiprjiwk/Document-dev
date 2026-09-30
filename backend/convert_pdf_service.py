import os
import io
import re
import datetime
import zipfile
import docx
import xml.etree.ElementTree as ET

def is_plan_report_filename(filename: str) -> bool:
    """
    Checks if filename contains 'plan-report' (case-insensitive).
    Files containing 'Plan-Report' are exempt from force-black text color conversion.
    """
    if not filename:
        return False
    return "plan-report" in filename.lower()

def compute_pdf_filename(filename: str) -> str:
    """
    Replaces .docx or .doc extension with .pdf.
    """
    base = os.path.splitext(filename)[0]
    return base + ".pdf"

def make_docx_text_black_in_memory(docx_bytes: bytes) -> bytes:
    """
    Pre-processes Word .docx bytes in memory using python-docx to change all run text colors to Black (#000000).
    """
    if not docx_bytes or not docx_bytes.startswith(b'PK\x03\x04'):
        return docx_bytes

    try:
        doc = docx.Document(io.BytesIO(docx_bytes))

        def _apply_black(p):
            for r in p.runs:
                try:
                    r.font.color.rgb = docx.shared.RGBColor(0, 0, 0)
                except Exception:
                    pass

        for p in doc.paragraphs:
            _apply_black(p)

        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _apply_black(p)

        for section in doc.sections:
            if section.header:
                for p in section.header.paragraphs:
                    _apply_black(p)
                for tbl in section.header.tables:
                    for row in tbl.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                _apply_black(p)
            if section.footer:
                for p in section.footer.paragraphs:
                    _apply_black(p)
                for tbl in section.footer.tables:
                    for row in tbl.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                _apply_black(p)

        out_io = io.BytesIO()
        doc.save(out_io)
        out_io.seek(0)
        return out_io.getvalue()
    except Exception:
        return docx_bytes

def process_batch_convert_pdf(
    file_items: list[dict],
    force_black_text: bool = True
) -> list[dict]:
    """
    Batch converts Word document files (.docx, .doc) to PDF (.pdf) format using 1 single persistent MS Word COM instance.
    
    Rule for force_black_text:
    - If force_black_text is True and filename does NOT contain 'Plan-Report', all text color is converted to Black before PDF export.
    - If filename contains 'Plan-Report', text color conversion is skipped (exempted) to preserve original styling.
    
    file_items structure:
    [
        {
            "filename": "XXXXX_01-00_IQ_Plan-Report_YYYY-MM-DD_en.docx",
            "bytes": b'...',
            "source_path": r"C:\...\file.docx" (optional)
        }, ...
    ]
    
    Returns list of dicts:
    [
        {
            "original_filename": "...",
            "pdf_filename": "...",
            "pdf_bytes": b'...',
            "is_plan_report_exempt": True/False,
            "status": "OK" / error_message,
            "source_path": "..."
        }, ...
    ]
    """
    results = []
    if not file_items:
        return results

    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()

        temp_dir = os.path.abspath("scratch")
        os.makedirs(temp_dir, exist_ok=True)

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0  # wdAlertsNone
        try:
            word.ScreenUpdating = False
        except Exception:
            pass

        try:
            for idx, item in enumerate(file_items):
                orig_filename = item["filename"]
                docx_bytes = item["bytes"]
                src_path = item.get("source_path")

                pdf_filename = compute_pdf_filename(orig_filename)
                is_exempt = is_plan_report_filename(orig_filename)
                should_make_black = force_black_text and not is_exempt

                # Pre-apply black text in memory if needed
                if should_make_black and docx_bytes and docx_bytes.startswith(b'PK\x03\x04'):
                    docx_bytes = make_docx_text_black_in_memory(docx_bytes)

                timestamp_str = f"{os.getpid()}_{int(datetime.datetime.now().timestamp() * 1000)}_{idx}"
                sub_temp_dir = os.path.join(temp_dir, f"pdf_work_{timestamp_str}")
                os.makedirs(sub_temp_dir, exist_ok=True)

                temp_in_path = os.path.join(sub_temp_dir, orig_filename)
                temp_pdf_path = os.path.join(sub_temp_dir, pdf_filename)

                with open(temp_in_path, "wb") as f:
                    f.write(docx_bytes)

                try:
                    doc = word.Documents.Open(
                        FileName=temp_in_path,
                        ConfirmConversions=False,
                        ReadOnly=True,
                        AddToRecentFiles=False
                    )

                    # Force black text color via Word COM API for non-Plan-Report files
                    if should_make_black:
                        try:
                            doc.Content.Font.Color = 0  # 0 = wdColorBlack
                            for story in doc.StoryRanges:
                                try:
                                    story.Font.Color = 0
                                except Exception:
                                    pass
                            for tbl in doc.Tables:
                                try:
                                    tbl.Range.Font.Color = 0
                                except Exception:
                                    pass
                        except Exception:
                            pass

                    # Export as PDF (17 = wdExportFormatPDF)
                    doc.ExportAsFixedFormat(
                        OutputFileName=temp_pdf_path,
                        ExportFormat=17,
                        OpenAfterExport=False,
                        OptimizeFor=0,  # 0 = wdExportOptimizeForPrint
                        Range=0,        # 0 = wdExportAllDocument
                        From=1,
                        To=1,
                        Item=0,         # 0 = wdExportDocumentContent
                        IncludeDocProps=True,
                        KeepIRM=True,
                        CreateBookmarks=1,  # 1 = wdExportCreateHeadingBookmarks
                        DocStructureTags=True,
                        BitmapMissingFonts=True,
                        UseISO19005_1=False
                    )

                    doc.Close(False)

                    if os.path.exists(temp_pdf_path):
                        with open(temp_pdf_path, "rb") as pdf_f:
                            pdf_bytes = pdf_f.read()
                        
                        results.append({
                            "original_filename": orig_filename,
                            "pdf_filename": pdf_filename,
                            "pdf_bytes": pdf_bytes,
                            "is_plan_report_exempt": is_exempt,
                            "status": "OK",
                            "source_path": src_path
                        })
                    else:
                        results.append({
                            "original_filename": orig_filename,
                            "pdf_filename": pdf_filename,
                            "pdf_bytes": None,
                            "is_plan_report_exempt": is_exempt,
                            "status": "Failed to generate PDF file.",
                            "source_path": src_path
                        })

                except Exception as ex:
                    results.append({
                        "original_filename": orig_filename,
                        "pdf_filename": pdf_filename,
                        "pdf_bytes": None,
                        "is_plan_report_exempt": is_exempt,
                        "status": f"Word COM Error: {str(ex)}",
                        "source_path": src_path
                    })

                # Cleanup sub temp directory
                try:
                    import shutil
                    shutil.rmtree(sub_temp_dir, ignore_errors=True)
                except Exception:
                    pass

        finally:
            try:
                word.Quit()
            except Exception:
                pass

    except Exception as g_ex:
        for item in file_items:
            fn = item["filename"]
            pdf_fn = compute_pdf_filename(fn)
            results.append({
                "original_filename": fn,
                "pdf_filename": pdf_fn,
                "pdf_bytes": None,
                "is_plan_report_exempt": is_plan_report_filename(fn),
                "status": f"COM Initializer Error: {str(g_ex)}",
                "source_path": item.get("source_path")
            })

    return results
