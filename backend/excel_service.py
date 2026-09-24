import io
import logging
import openpyxl
from typing import Dict, Any, Union

# Set up logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger("ExcelExportService")

class ExcelExportService:
    """
    Service responsible for loading an Excel workbook template, populating cells
    using mapped OCR keys, and exporting the completed file without destroying formulas or formatting.
    """
    
    def __init__(self, default_sheet_name: str = "Blank form"):
        self.default_sheet_name = default_sheet_name

    def export_excel(
        self,
        template_bytes: bytes,
        ocr_results: Dict[str, Any]
    ) -> bytes:
        """
        Populates Excel template cells based on ocr_results.
        Supports sheet-qualified cell addresses (e.g. 'F1!C4' or 'F2!F12') or simple cells ('C4').
        
        Args:
            template_bytes (bytes): The raw bytes of the template Excel workbook.
            ocr_results (Dict[str, Any]): Dictionary of fields containing 'Excel Cell', 'Corrected Value', etc.
            
        Returns:
            bytes: The completed Excel workbook bytes ready for export.
        """
        logger.info("Initializing multi-sheet Excel template writeback process...")
        
        try:
            wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
            write_count = 0
            
            for field_name, item in ocr_results.items():
                cell_raw = item.get("Excel Cell") or item.get("cell_addr")
                if not cell_raw:
                    continue
                    
                val = item.get("Corrected Value")
                if val is None:
                    val = item.get("Detected Value", "")
                    
                target_sheet = self.default_sheet_name
                cell_coord = str(cell_raw).strip()
                
                # Check for Sheet!Cell format
                if "!" in cell_coord:
                    parts = cell_coord.split("!", 1)
                    target_sheet = parts[0].strip()
                    cell_coord = parts[1].strip()
                elif item.get("Sheet"):
                    target_sheet = str(item["Sheet"]).strip()
                    
                if target_sheet in wb.sheetnames:
                    ws = wb[target_sheet]
                else:
                    ws = wb.active
                    logger.warning(f"Target sheet '{target_sheet}' not found. Defaulting to '{ws.title}'")
                    
                val_str = str(val).strip() if val is not None else ""
                
                if val_str == "":
                    ws[cell_coord] = None
                elif val_str.startswith('0') and len(val_str) > 1 and not val_str.startswith('0.'):
                    ws[cell_coord] = val_str
                else:
                    try:
                        if '.' in val_str:
                            ws[cell_coord] = float(val_str)
                        else:
                            ws[cell_coord] = int(val_str)
                    except ValueError:
                        ws[cell_coord] = val
                        
                write_count += 1
                logger.info(f"Wrote to [{ws.title}] {cell_coord} -> '{ws[cell_coord].value}' for field '{field_name}'")
                
            out_bytes = io.BytesIO()
            wb.save(out_bytes)
            logger.info(f"Completed multi-sheet Excel export. Total cells updated: {write_count}")
            return out_bytes.getvalue()
            
        except Exception as e:
            logger.error(f"Error during excel multi-sheet export: {e}")
            raise ValueError(f"Failed to export Excel workbook: {str(e)}")
