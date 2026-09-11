import os
import io
import fitz  # PyMuPDF
from pypdf import PdfWriter, PdfReader
import pandas as pd
import zipfile

def parse_and_transform_bom(bom_file_bytes, cert_type, matched_files_set=None, missing_files_set=None, is_manual=False):
    """
    Parses and transforms an uploaded BOM Excel file based on certificate type using openpyxl.
    Supports both Original BOM (auto-transformed via macro) and Manual / Pre-transformed BOM.
    Applies Bold font for Level 2, Red fill for macro condition fails,
    Green fill for matched PDFs, and Yellow fill for missing PDFs.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    import pandas as pd
    
    try:
        wb = openpyxl.load_workbook(io.BytesIO(bom_file_bytes))
        if 'raw_data' in wb.sheetnames:
            ws = wb['raw_data']
        else:
            ws = wb.active
            
        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        green_fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        orange_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        bold_font = Font(bold=True)

        if is_manual:
            # Manual / Pre-transformed BOM: skip structural deletions, apply formatting & fills
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                if colA == "2":
                    for c in range(1, 16):
                        ws.cell(row=r, column=c).font = bold_font

            for r in range(2, ws.max_row + 1):
                colC = str(ws.cell(row=r, column=3).value or "").strip()
                if colC.endswith(".0"): colC = colC[:-2]
                colC_clean = colC.lower()

                is_red = False
                if cert_type == "WAZ and FAD":
                    valL = str(ws.cell(row=r, column=12).value or "").strip()
                    if valL == "": is_red = True
                elif cert_type == "OZ":
                    valJ = str(ws.cell(row=r, column=10).value or "").strip()
                    valK = str(ws.cell(row=r, column=11).value or "").strip()
                    valL = str(ws.cell(row=r, column=12).value or "").strip()
                    if valJ == "" or valK == "" or valL == "": is_red = True
                elif cert_type in ["SZ", "OMP"]:
                    valK = str(ws.cell(row=r, column=11).value or "").strip()
                    valL = str(ws.cell(row=r, column=12).value or "").strip()
                    if valK == "" or valL == "": is_red = True

                if is_red:
                    for c in range(1, 14):
                        ws.cell(row=r, column=c).fill = red_fill

                if matched_files_set and colC_clean in matched_files_set:
                    ws.cell(row=r, column=3).fill = green_fill
                elif missing_files_set and colC_clean in missing_files_set:
                    ws.cell(row=r, column=3).fill = yellow_fill

        elif cert_type == "WAZ and FAD":
            # Step 1: Delete rows above "Product contact parts" in col D
            product_row = None
            for r in range(2, ws.max_row + 1):
                val = str(ws.cell(row=r, column=4).value or "")
                if val == "Product contact parts":
                    product_row = r
                    break
            if product_row and product_row > 1:
                ws.delete_rows(2, amount=product_row - 1)

            # Step 2: Bold rows where col A == "2"
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                if colA == "2":
                    for c in range(1, 18):
                        ws.cell(row=r, column=c).font = bold_font

            # Step 3: Copy col K (11) to Q (17)
            for r in range(1, ws.max_row + 1):
                ws.cell(row=r, column=17).value = ws.cell(row=r, column=11).value

            # Step 4: Delete cols K to O (11 to 15)
            ws.delete_cols(11, amount=5)

            # Step 5: Delete rows where col D contains "Electrical parts"
            for r in range(ws.max_row, 1, -1):
                valD = str(ws.cell(row=r, column=4).value or "")
                if "Electrical parts" in valD:
                    ws.delete_rows(r, 1)

            # Step 6: Filter rows based on conditions
            for r in range(ws.max_row, 1, -1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                colB = str(ws.cell(row=r, column=2).value or "").upper()
                colI = str(ws.cell(row=r, column=9).value or "").upper()
                colJ = str(ws.cell(row=r, column=10).value or "").upper()
                
                keep_row = False
                if colA == "2":
                    next_colA = ""
                    if r + 1 <= ws.max_row:
                        next_colA = str(ws.cell(row=r+1, column=1).value or "").strip()
                        if next_colA.endswith(".0"): next_colA = next_colA[:-2]
                    if next_colA not in ["1", "2"]:
                        keep_row = True
                        
                if "3.1" in colI or "FDA" in colJ:
                    keep_row = True
                    
                if "PRODUKTBERÜHREND" in colB or "PRODUKTBERUEHREND" in colB:
                    keep_row = True
                    
                if not keep_row:
                    ws.delete_rows(r, 1)

            # Step 7: Generate result text in col M (col 13)
            seq_number = 1
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                colC = str(ws.cell(row=r, column=3).value or "").strip()
                if colC.endswith(".0"): colC = colC[:-2]
                
                if colA == "2":
                    # Scan child rows of this group
                    has_31 = False
                    has_fda = False
                    
                    self_I = str(ws.cell(row=r, column=9).value or "").upper()
                    self_J = str(ws.cell(row=r, column=10).value or "").upper()
                    if "3.1" in self_I: has_31 = True
                    if "FDA" in self_J: has_fda = True
                    
                    next_r = r + 1
                    while next_r <= ws.max_row:
                        next_colA = str(ws.cell(row=next_r, column=1).value or "").strip()
                        if next_colA.endswith(".0"): next_colA = next_colA[:-2]
                        
                        if next_colA in ["1", "2"]:
                            break
                            
                        child_I = str(ws.cell(row=next_r, column=9).value or "").upper()
                        child_J = str(ws.cell(row=next_r, column=10).value or "").upper()
                        
                        if "3.1" in child_I:
                            has_31 = True
                        if "FDA" in child_J:
                            has_fda = True
                            
                        next_r += 1
                        
                    res = f"{seq_number:03d}_{colC}_"
                    if has_31 and has_fda:
                        res += "WAZ_FDA"
                    elif has_31 and not has_fda:
                        res += "WAZ"
                    elif has_fda and not has_31:
                        res += "FDA"
                    else:
                        res += "FDA"  # Default if neither is found
                        
                    ws.cell(row=r, column=13).value = res
                    seq_number += 1
                else:
                    ws.cell(row=r, column=13).value = colC

            # Step 8: Delete trailing rows if col I or J doesn't contain "3.1" or "FDA"
            while ws.max_row >= 2:
                colI = str(ws.cell(row=ws.max_row, column=9).value or "").upper()
                colJ = str(ws.cell(row=ws.max_row, column=10).value or "").upper()
                if "3.1" in colI or "FDA" in colJ:
                    break
                else:
                    ws.delete_rows(ws.max_row, 1)

            # Step 9 & Status Coloring: Highlight A:O
            for r in range(2, ws.max_row + 1):
                valL = str(ws.cell(row=r, column=12).value or "").strip()
                colC = str(ws.cell(row=r, column=3).value or "").strip()
                if colC.endswith(".0"): colC = colC[:-2]
                colC_clean = colC.lower()
                
                if valL == "":
                    # Original Macro rule: Red for blank col L
                    for c in range(1, 16):
                        ws.cell(row=r, column=c).fill = red_fill
                        
                # Green/Yellow override on col C (Component number)
                if matched_files_set and colC_clean in matched_files_set:
                    ws.cell(row=r, column=3).fill = green_fill
                elif missing_files_set and colC_clean in missing_files_set:
                    ws.cell(row=r, column=3).fill = yellow_fill

        elif cert_type == "OZ":
            # Step 2: Bold rows where col A == "2"
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                if colA == "2":
                    for c in range(1, 16):
                        ws.cell(row=r, column=c).font = bold_font

            # Step 3: Highlight "product contact parts" in col C and delete previous rows
            found_prod = None
            for r in range(2, ws.max_row + 1):
                valC = str(ws.cell(row=r, column=3).value or "")
                if "product contact parts" in valC.lower():
                    ws.cell(row=r, column=3).fill = orange_fill
                    found_prod = r
                    break
            if found_prod and found_prod >= 2:
                ws.delete_rows(2, amount=found_prod - 1)

            # Step 4: Delete from "Electrical components options" in col D to end
            found_opt = None
            for r in range(2, ws.max_row + 1):
                valD = str(ws.cell(row=r, column=4).value or "")
                if "Electrical components options" in valD.lower():
                    found_opt = r
                    break
            if found_opt:
                ws.delete_rows(found_opt, amount=ws.max_row - found_opt + 1)

            # Step 5: Delete rows based on conditions
            for r in range(ws.max_row, 1, -1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                colB = str(ws.cell(row=r, column=2).value or "").upper()
                valM = str(ws.cell(row=r, column=13).value or "").upper()
                valO = str(ws.cell(row=r, column=15).value or "").upper()
                
                keep_row = False
                if colA == "2":
                    keep_row = True
                if "ZERTIFIKAT FÜR E-POLIEREN" in valM or "PASSIVIERUNGSZERTIFIKAT" in valO:
                    keep_row = True
                if "PRODUKTBERÜHREND" in colB or "PRODUKTBERUEHREND" in colB:
                    keep_row = True
                    
                if not keep_row:
                    ws.delete_rows(r, 1)

            # Step 6: Delete cols J, L, N (cols 14, 12, 10)
            ws.delete_cols(14, 1)
            ws.delete_cols(12, 1)
            ws.delete_cols(10, 1)

            # Step 8: Fill col M (13)
            seq_num = 1
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                rowC = str(ws.cell(row=r, column=3).value or "").strip()
                if rowC.endswith(".0"): rowC = rowC[:-2]
                if rowC:
                    if colA == "2":
                        ws.cell(row=r, column=13).value = f"{seq_num:03d}_{rowC}_OZ"
                        seq_num += 1
                    else:
                        ws.cell(row=r, column=13).value = rowC

            # Step 9 & Status Coloring: Highlight A:M
            for r in range(2, ws.max_row + 1):
                colJ = str(ws.cell(row=r, column=10).value or "").strip()
                colK = str(ws.cell(row=r, column=11).value or "").strip()
                colL = str(ws.cell(row=r, column=12).value or "").strip()
                colC = str(ws.cell(row=r, column=3).value or "").strip()
                if colC.endswith(".0"): colC = colC[:-2]
                colC_clean = colC.lower()
                
                if colJ == "" or colK == "" or colL == "":
                    for c in range(1, 14):
                        ws.cell(row=r, column=c).fill = red_fill
                
                # Green/Yellow override on col C (Component number)
                if matched_files_set and colC_clean in matched_files_set:
                    ws.cell(row=r, column=3).fill = green_fill
                elif missing_files_set and colC_clean in missing_files_set:
                    ws.cell(row=r, column=3).fill = yellow_fill

        elif cert_type == "SZ":
            # Step 2: Bold rows where col A == "2"
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                if colA == "2":
                    for c in range(1, 16):
                        ws.cell(row=r, column=c).font = bold_font

            # Step 3: Find "product contact parts" in col C. Highlight RGB(255,192,0). Delete 2:foundProductRow.
            found_prod = None
            for r in range(2, ws.max_row + 1):
                valC = str(ws.cell(row=r, column=3).value or "")
                if "product contact parts" in valC.lower():
                    ws.cell(row=r, column=3).fill = orange_fill
                    found_prod = r
                    break
            if found_prod and found_prod >= 2:
                ws.delete_rows(2, amount=found_prod - 1)

            # Step 4: Delete from "Electrical components options" to end
            found_opt = None
            for r in range(2, ws.max_row + 1):
                valD = str(ws.cell(row=r, column=4).value or "")
                if "Electrical components options" in valD.lower():
                    found_opt = r
                    break
            if found_opt:
                ws.delete_rows(found_opt, amount=ws.max_row - found_opt + 1)

            # Step 5: Delete rows based on conditions
            for r in range(ws.max_row, 1, -1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                colB = str(ws.cell(row=r, column=2).value or "").upper()
                valN = str(ws.cell(row=r, column=14).value or "").upper()
                
                keep_row = False
                if colA == "2":
                    keep_row = True
                if "SCHWEIßZERTIFIKAT" in valN:
                    keep_row = True
                if "PRODUKTBERÜHREND" in colB or "PRODUKTBERUEHREND" in colB:
                    keep_row = True
                    
                if not keep_row:
                    ws.delete_rows(r, 1)

            # Step 6: Delete cols J, L, M, O (15, 13, 12, 10)
            ws.delete_cols(15, 1)
            ws.delete_cols(13, 1)
            ws.delete_cols(12, 1)
            ws.delete_cols(10, 1)

            # Step 7: Insert empty col at J (10)
            ws.insert_cols(10, 1)

            # Step 8: Fill col M (13)
            seq_num = 1
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                rowC = str(ws.cell(row=r, column=3).value or "").strip()
                if rowC.endswith(".0"): rowC = rowC[:-2]
                if rowC:
                    if colA == "2":
                        ws.cell(row=r, column=13).value = f"{seq_num:03d}_{rowC}_SZ"
                        seq_num += 1
                    else:
                        ws.cell(row=r, column=13).value = rowC

            # Step 9 & Status Coloring: Highlight A:M
            for r in range(2, ws.max_row + 1):
                colK = str(ws.cell(row=r, column=11).value or "").strip()
                colL = str(ws.cell(row=r, column=12).value or "").strip()
                colC = str(ws.cell(row=r, column=3).value or "").strip()
                if colC.endswith(".0"): colC = colC[:-2]
                colC_clean = colC.lower()
                
                if colK == "" or colL == "":
                    for c in range(1, 14):
                        ws.cell(row=r, column=c).fill = red_fill
                
                # Green/Yellow override on col C (Component number)
                if matched_files_set and colC_clean in matched_files_set:
                    ws.cell(row=r, column=3).fill = green_fill
                elif missing_files_set and colC_clean in missing_files_set:
                    ws.cell(row=r, column=3).fill = yellow_fill

        elif cert_type == "OMP":
            # Step 2: Bold rows where col A == "2"
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                if colA == "2":
                    for c in range(1, 16):
                        ws.cell(row=r, column=c).font = bold_font

            # Step 3: Find "product contact parts" in col C. Highlight RGB(255,192,0). Delete 2:foundProductRow.
            found_prod = None
            for r in range(2, ws.max_row + 1):
                valC = str(ws.cell(row=r, column=3).value or "")
                if "product contact parts" in valC.lower():
                    ws.cell(row=r, column=3).fill = orange_fill
                    found_prod = r
                    break
            if found_prod and found_prod >= 2:
                ws.delete_rows(2, amount=found_prod - 1)

            # Step 4: Delete from "Electrical components options" to end
            found_opt = None
            for r in range(2, ws.max_row + 1):
                valD = str(ws.cell(row=r, column=4).value or "")
                if "Electrical components options" in valD.lower():
                    found_opt = r
                    break
            if found_opt:
                ws.delete_rows(found_opt, amount=ws.max_row - found_opt + 1)

            # Step 5: Delete rows based on conditions
            for r in range(ws.max_row, 1, -1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                colB = str(ws.cell(row=r, column=2).value or "").upper()
                valL = str(ws.cell(row=r, column=12).value or "").upper()
                
                keep_row = False
                if colA == "2":
                    keep_row = True
                if "OBERFLÄCHENMESSPROTOKOLL" in valL:
                    keep_row = True
                if "PRODUKTBERÜHREND" in colB or "PRODUKTBERUEHREND" in colB:
                    keep_row = True
                    
                if not keep_row:
                    ws.delete_rows(r, 1)

            # Step 6: Delete cols J, M, N, O (15, 14, 13, 10)
            ws.delete_cols(15, 1)
            ws.delete_cols(14, 1)
            ws.delete_cols(13, 1)
            ws.delete_cols(10, 1)

            # Step 7: Insert empty col at J (10)
            ws.insert_cols(10, 1)

            # Step 8: Fill col M (13)
            seq_num = 1
            for r in range(2, ws.max_row + 1):
                colA = str(ws.cell(row=r, column=1).value or "").strip()
                if colA.endswith(".0"): colA = colA[:-2]
                rowC = str(ws.cell(row=r, column=3).value or "").strip()
                if rowC.endswith(".0"): rowC = rowC[:-2]
                if rowC:
                    if colA == "2":
                        ws.cell(row=r, column=13).value = f"{seq_num:03d}_{rowC}_OMP"
                        seq_num += 1
                    else:
                        ws.cell(row=r, column=13).value = rowC

            # Step 9 & Status Coloring: Highlight A:M
            for r in range(2, ws.max_row + 1):
                colK = str(ws.cell(row=r, column=11).value or "").strip()
                colL = str(ws.cell(row=r, column=12).value or "").strip()
                colC = str(ws.cell(row=r, column=3).value or "").strip()
                if colC.endswith(".0"): colC = colC[:-2]
                colC_clean = colC.lower()
                
                if colK == "" or colL == "":
                    for c in range(1, 14):
                        ws.cell(row=r, column=c).fill = red_fill
                
                # Green/Yellow override on col C (Component number)
                if matched_files_set and colC_clean in matched_files_set:
                    ws.cell(row=r, column=3).fill = green_fill
                elif missing_files_set and colC_clean in missing_files_set:
                    ws.cell(row=r, column=3).fill = yellow_fill

        # Extract records for PDF grouping
        records = []
        current_group = "Unknown_Group"
        for r in range(2, ws.max_row + 1):
            colA = str(ws.cell(row=r, column=1).value or "").strip()
            if colA.endswith(".0"): colA = colA[:-2]
            colC = str(ws.cell(row=r, column=3).value or "").strip()
            if colC.endswith(".0"): colC = colC[:-2]
            colM = str(ws.cell(row=r, column=13).value or "").strip()
            
            if colA == "2":
                current_group = colM
                records.append({"File Name": "", "Group": current_group})
            else:
                if colC:
                    records.append({"File Name": f"{colC}.pdf", "Group": current_group})

        # Convert to DataFrame for Streamlit preview with unique column names
        data = list(ws.values)
        if data:
            raw_header = data[0]
            header = []
            seen = {}
            for i, col in enumerate(raw_header):
                c_str = str(col).strip() if col is not None else f"Column_{i+1}"
                if c_str in seen:
                    seen[c_str] += 1
                    header.append(f"{c_str}_{seen[c_str]}")
                else:
                    seen[c_str] = 0
                    header.append(c_str)
            rows = data[1:]
            df_preview = pd.DataFrame(rows, columns=header)
        else:
            df_preview = pd.DataFrame()

        out_buf = io.BytesIO()
        wb.save(out_buf)
        return True, records, df_preview, out_buf.getvalue()
    except Exception as e:
        return False, str(e), None, None

def parse_bom(bom_file_bytes, cert_type):
    success, records, df_preview, excel_bytes = parse_and_transform_bom(bom_file_bytes, cert_type)
    return success, records, df_preview

def style_bom_dataframe(df, cert_type, matched_files_set=None, missing_files_set=None):
    """
    Applies CSS styling to the BOM DataFrame for live Streamlit preview:
    1. Bold text for Level 2 rows.
    2. Red background for Macro condition fails.
    3. Green background for Add Text success (matched PDF).
    4. Yellow background for Add Text fail (missing PDF).
    """
    if df is None or df.empty:
        return df

    def highlight_row(row):
        colA = str(row.iloc[0]).strip() if len(row) > 0 and pd.notna(row.iloc[0]) else ""
        if colA.endswith(".0"): colA = colA[:-2]
        
        colC = str(row.iloc[2]).strip() if len(row) > 2 and pd.notna(row.iloc[2]) else ""
        if colC.endswith(".0"): colC = colC[:-2]
        colC_clean = colC.lower()

        styles = [''] * len(row)

        # 1. Level 2 -> Bold font
        if colA == "2":
            styles = ['font-weight: bold'] * len(row)

        # 2. Red bar condition (Macro rule fail)
        is_red = False
        if cert_type == "WAZ and FAD":
            valL = str(row.iloc[11]).strip() if len(row) > 11 and pd.notna(row.iloc[11]) else ""
            if valL == "" or valL == "None":
                is_red = True
        elif cert_type == "OZ":
            valJ = str(row.iloc[9]).strip() if len(row) > 9 and pd.notna(row.iloc[9]) else ""
            valK = str(row.iloc[10]).strip() if len(row) > 10 and pd.notna(row.iloc[10]) else ""
            valL = str(row.iloc[11]).strip() if len(row) > 11 and pd.notna(row.iloc[11]) else ""
            if valJ in ["", "None"] or valK in ["", "None"] or valL in ["", "None"]:
                is_red = True
        elif cert_type in ["SZ", "OMP"]:
            valK = str(row.iloc[10]).strip() if len(row) > 10 and pd.notna(row.iloc[10]) else ""
            valL = str(row.iloc[11]).strip() if len(row) > 11 and pd.notna(row.iloc[11]) else ""
            if valK in ["", "None"] or valL in ["", "None"]:
                is_red = True

        if is_red:
            styles = [s + ('; ' if s else '') + 'background-color: #ff4d4d; color: white' for s in styles]

        # For Green and Yellow, ONLY apply to Component number cell (index 2), overriding Red row background
        if matched_files_set and colC_clean in matched_files_set:
            styles[2] = 'background-color: #4caf50; color: white; font-weight: bold'
        elif missing_files_set and colC_clean in missing_files_set:
            styles[2] = 'background-color: #ffca28; color: black; font-weight: bold'

        return styles

    return df.style.apply(highlight_row, axis=1)

def separate_and_match_files(uploaded_files, bom_records):
    groups = {}
    missing_files = []
    matched_files = {}
    used_files = set()
    
    uploaded_map = {}
    for f in uploaded_files:
        base_name = os.path.splitext(f.name)[0].lower()
        uploaded_map[base_name] = f
        
    for record in bom_records:
        fname_base = str(record.get('File Name', '')).strip().lower()
        if fname_base.endswith('.pdf'): fname_base = fname_base[:-4]
        group = str(record.get('Group', 'Unknown'))
        
        if group not in groups:
            groups[group] = []
            
        if fname_base:
            if fname_base in uploaded_map:
                matched_f = uploaded_map[fname_base]
                groups[group].append(matched_f.name)
                matched_files[matched_f.name] = matched_f
                used_files.add(matched_f.name)
            else:
                missing_files.append({"File Name": record.get('File Name', ''), "Group": group})
            
    return groups, matched_files, missing_files, used_files

def insert_text_visual(page, p_vis, text, fontsize=14, align="left"):
    text_len = fitz.get_text_length(text, fontname="helv", fontsize=fontsize)
    if align == "center":
        p_vis = fitz.Point(p_vis.x - text_len / 2, p_vis.y)
    elif align == "right":
        p_vis = fitz.Point(p_vis.x - text_len, p_vis.y)
        
    p_unrot = p_vis * ~page.rotation_matrix
    m = fitz.Matrix(page.rotation)
    page.insert_text(p_unrot, text, fontsize=fontsize, fontname="helv", color=(0, 0, 0), morph=(p_unrot, m))

def add_cert_part_text(pdf_bytes, certificate_no, part_no):
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page in doc:
            # Left (Certificate No)
            insert_text_visual(page, fitz.Point(30, 30), f"Certificate No.: {certificate_no}", fontsize=14, align="left")
            
            # Center (Part No)
            insert_text_visual(page, fitz.Point(page.rect.width / 2, 30), f"Part No.: {part_no}", fontsize=14, align="center")
            
        output_pdf = io.BytesIO()
        doc.save(output_pdf)
        doc.close()
        return output_pdf.getvalue()
    except Exception as e:
        print(f"Error adding cert/part text: {e}")
        return pdf_bytes

def add_page_numbers(pdf_bytes):
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        total_pages = len(doc)
        for i, page in enumerate(doc):
            text_page = f"Page {i + 1} of {total_pages}"
            
            # Right (Page X of Y)
            insert_text_visual(page, fitz.Point(page.rect.width - 30, 30), text_page, fontsize=14, align="right")
            
        output_pdf = io.BytesIO()
        doc.save(output_pdf)
        doc.close()
        return output_pdf.getvalue()
    except Exception as e:
        print(f"Error adding page numbers: {e}")
        return pdf_bytes

def merge_pdfs_in_group(file_objects_bytes):
    writer = PdfWriter()
    for pdf_bytes in file_objects_bytes:
        try:
            reader = PdfReader(io.BytesIO(pdf_bytes))
            for page in reader.pages:
                writer.add_page(page)
        except Exception as e:
            print(f"Error merging file: {e}")
            
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()

def process_iwk_certificates(uploaded_files, bom_records, cert_type, bom_file_bytes=None, is_manual=False):
    import pandas as pd
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    
    groups, matched_files, missing_files, used_files = separate_and_match_files(uploaded_files, bom_records)
    
    zip_buffer = io.BytesIO()
    final_pdfs = {}
    status_data = []
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for group_name, file_names in groups.items():
            if not file_names:
                # Create an empty folder for the group so the sequence doesn't skip
                zip_info = zipfile.ZipInfo(f"{group_name}/")
                zip_file.writestr(zip_info, b'')
                continue
                
            group_pdf_bytes = []
            # Extract certificate_no from group name (e.g. "001" from "001_2622364_WAZ")
            certificate_no = group_name.split('_')[0] if '_' in group_name else group_name
            
            for fname in file_names:
                f_obj = matched_files[fname]
                pos = f_obj.tell()
                f_obj.seek(0)
                f_bytes = f_obj.read()
                f_obj.seek(pos)
                
                if fname.lower().endswith('.pdf'):
                    part_no = os.path.splitext(fname)[0]
                    # Code 1: Add Cert No and Part No to each file
                    modified_bytes = add_cert_part_text(f_bytes, certificate_no, part_no)
                    group_pdf_bytes.append(modified_bytes)
                    status_data.append({"Part No": part_no, "Status": "Success", "Certificate No": certificate_no})
                else:
                    zip_file.writestr(f"{group_name}/{fname}", f_bytes)
                    
            if group_pdf_bytes:
                # Merge files
                merged_bytes = merge_pdfs_in_group(group_pdf_bytes)
                # Code 2: Add Page Numbers to the merged file
                final_merged_bytes = add_page_numbers(merged_bytes)
                zip_file.writestr(f"{group_name}/{group_name}.pdf", final_merged_bytes)
                final_pdfs[f"{group_name}.pdf"] = final_merged_bytes
                
            
    # Generate 3-Sheet Backup Excel Status
    excel_status_bytes = None
    if bom_file_bytes:
        try:
            matched_set = {os.path.splitext(f)[0].lower() for f in matched_files.keys()}
            missing_set = {str(item.get("File Name", "")).replace(".pdf", "").lower() for item in missing_files}
            
            # Load original workbook for Sheet 1
            wb_orig = openpyxl.load_workbook(io.BytesIO(bom_file_bytes))
            ws_orig = wb_orig['raw_data'] if 'raw_data' in wb_orig.sheetnames else wb_orig.active
            
            # Generate transformed workbook for Sheet 2
            _, _, _, trans_bytes = parse_and_transform_bom(bom_file_bytes, cert_type, is_manual=is_manual)
            wb_trans = openpyxl.load_workbook(io.BytesIO(trans_bytes))
            ws_trans = wb_trans.active
            
            # Generate processed workbook for Sheet 3
            _, _, _, proc_bytes = parse_and_transform_bom(
                bom_file_bytes, cert_type, matched_files_set=matched_set, missing_files_set=missing_set, is_manual=is_manual
            )
            wb_proc = openpyxl.load_workbook(io.BytesIO(proc_bytes))
            ws_proc = wb_proc.active
            
            # Create master 3-Sheet workbook
            master_wb = openpyxl.Workbook()
            master_wb.remove(master_wb.active) # Remove default sheet
            
            def copy_sheet_data(src_ws, target_wb, title):
                dest_ws = target_wb.create_sheet(title=title)
                for row in src_ws.iter_rows():
                    for cell in row:
                        new_cell = dest_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                        if cell.has_style:
                            if cell.font:
                                new_cell.font = Font(
                                    name=cell.font.name,
                                    size=cell.font.size,
                                    bold=cell.font.bold,
                                    italic=cell.font.italic,
                                    color=cell.font.color
                                )
                            if cell.fill and cell.fill.fill_type:
                                new_cell.fill = PatternFill(
                                    fill_type=cell.fill.fill_type,
                                    start_color=cell.fill.start_color,
                                    end_color=cell.fill.end_color
                                )

            copy_sheet_data(ws_orig, master_wb, "1_Original_BOM")
            copy_sheet_data(ws_trans, master_wb, "2_Transformed_BOM")
            copy_sheet_data(ws_proc, master_wb, "3_Process_Status")
            
            out_buf = io.BytesIO()
            master_wb.save(out_buf)
            excel_status_bytes = out_buf.getvalue()
        except Exception as e:
            print(f"Error generating 3-sheet Excel status: {e}")
                
    return zip_buffer.getvalue(), excel_status_bytes, final_pdfs
