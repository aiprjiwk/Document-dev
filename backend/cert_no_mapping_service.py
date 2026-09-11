import os
import re
import io
import zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill
import pandas as pd
import fitz  # PyMuPDF
from PIL import Image
import logging

logger = logging.getLogger("CertNoMappingService")
logging.basicConfig(level=logging.INFO)

# Optional OCR imports
try:
    from winocr import recognize_pil_sync as winocr_recognize
    WINOCR_AVAILABLE = True
except Exception:
    winocr_recognize = None
    WINOCR_AVAILABLE = False


def extract_cert_info_from_pdf(pdf_bytes: bytes, filename: str) -> dict:
    """
    Extracts Serial No. and Certificate No. from a PDF file using text parsing and OCR fallbacks.
    Returns a dict with 'serial_no', 'cert_no', and 'status'.
    """
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        full_text = ""

        # Step 1: Extract embedded text via PyMuPDF
        for page in doc:
            full_text += page.get_text() + "\n"

        # Step 2: If text is empty or insufficient, use OCR (WinOCR)
        if len(full_text.strip()) < 20 and WINOCR_AVAILABLE and len(doc) > 0:
            logger.info(f"Running OCR on {filename}...")
            full_text = ""
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                pix = page.get_pixmap(dpi=200)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                res = winocr_recognize(img, "en")
                if res and "lines" in res:
                    page_text = " \n ".join([l["text"] for l in res["lines"]])
                    full_text += page_text + "\n"
        
        doc.close()

        # Step 3: Parse Serial No. and Certificate No. using regex patterns
        serial_no = None
        cert_no = None

        # Serial No regex matching (e.g. 56021-2532979-40-1 or 56021-2518166-1670-1)
        serial_match = re.findall(r'\b\d{5}-\d+-\d+-\d+\b', full_text)
        if serial_match:
            serial_no = serial_match[0].strip()
        else:
            # Fallback pattern for serial no
            sn_regex = re.search(r'Serial\s*No\.?\s*[:\s]*([A-Za-z0-9\-_]+)', full_text, re.IGNORECASE)
            if sn_regex:
                serial_no = sn_regex.group(1).strip()

        # Certificate No regex matching (e.g. PI-1407003/26 or PO-2908001/26)
        cert_match = re.findall(r'\b[A-Za-z]{2,4}-\d+/\d+\b', full_text)
        if cert_match:
            cert_no = cert_match[0].strip()
        else:
            # Fallback pattern for cert no
            cn_regex = re.search(r'Certificate\s*No\.?\s*[:\s]*([A-Za-z0-9\-/_]+)', full_text, re.IGNORECASE)
            if cn_regex:
                cert_no = cn_regex.group(1).strip()

        status = "Success" if (serial_no and cert_no) else ("Partial Extraction" if (serial_no or cert_no) else "Extraction Failed")

        return {
            "filename": filename,
            "serial_no": serial_no,
            "cert_no": cert_no,
            "status": status,
            "raw_text": full_text
        }
    except Exception as e:
        logger.error(f"Error extracting data from PDF {filename}: {e}")
        return {
            "filename": filename,
            "serial_no": None,
            "cert_no": None,
            "status": f"Error: {str(e)}",
            "raw_text": ""
        }


def sanitize_filename(filename: str) -> str:
    """
    Sanitizes a string to be safe for Windows file paths by replacing invalid characters.
    """
    invalid_chars = r'[\\/*?:"<>|]'
    sanitized = re.sub(invalid_chars, "_", filename)
    return sanitized.strip()


def process_certificate_no_mapping(excel_bytes: bytes, excel_filename: str, pdf_files_list: list) -> dict:
    """
    Main workflow function for Certificate No. Mapping Tool:
    1. Parse uploaded PDFs to extract Serial No. and Certificate No.
    2. Read Excel file (Step 2 layout), locate 'Serial No.' and 'Certificate No.' columns dynamically.
    3. Update Excel cells with extracted Certificate No. (Black font, non-bold/normal weight) where Serial No. matches.
    4. Highlight unmatched ENTIRE ROWS in Excel with Yellow fill (FFFF00).
    5. Inject IWK Logo (57.6 pt width) into Excel Left Header.
    6. Rename matched PDFs based on sequence number (001_, 002_) + Certificate No.
    7. Package updated Excel + ONLY matched PDFs into downloadable ZIP.
    8. Return detailed summary results.

    pdf_files_list element: dict with {"filename": str, "bytes": bytes}
    """
    logger.info(f"Processing Certificate No Mapping for {len(pdf_files_list)} PDFs against {excel_filename}")

    # Yellow fill pattern for unmatched items in Excel
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    # 1. Process all PDFs
    pdf_results = []
    serial_to_pdf_info = {}

    for pdf_item in pdf_files_list:
        p_filename = pdf_item["filename"]
        p_bytes = pdf_item["bytes"]
        info = extract_cert_info_from_pdf(p_bytes, p_filename)
        info["bytes"] = p_bytes
        pdf_results.append(info)

        if info["serial_no"]:
            serial_to_pdf_info[info["serial_no"].lower().strip()] = info

    # 2. Process Excel using openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    if 'Final Layout' in wb.sheetnames:
        ws = wb['Final Layout']
    else:
        ws = wb.active

    # Find headers dynamically in row 1
    serial_col_idx = None
    cert_col_idx = None
    no_col_idx = 1  # Column A / 1 is "No."

    for col in range(1, ws.max_column + 1):
        val = str(ws.cell(row=1, column=col).value or "").strip()
        if val.lower() == "no." or val.lower() == "no":
            no_col_idx = col
        elif val.lower() == "serial no." or val.lower() == "serial no":
            serial_col_idx = col
        elif val.lower() == "certificate no." or val.lower() == "certificate no":
            cert_col_idx = col

    if not serial_col_idx or not cert_col_idx:
        # Fallback to column 9 (Serial No.) and 8 (Certificate No.) if headers not explicitly matched
        serial_col_idx = serial_col_idx or 9
        cert_col_idx = cert_col_idx or 8
        logger.warning(f"Using column fallback: Serial No col={serial_col_idx}, Cert No col={cert_col_idx}")

    # 3. Match and Update Excel cells
    rows_updated = 0
    mapping_records = []

    for r in range(2, ws.max_row + 1):
        cell_serial_val = str(ws.cell(row=r, column=serial_col_idx).value or "").strip()
        cell_serial_clean = cell_serial_val.lower()

        matched_pdf = None
        if cell_serial_clean in serial_to_pdf_info:
            matched_pdf = serial_to_pdf_info[cell_serial_clean]

        if matched_pdf and matched_pdf["cert_no"]:
            cert_val = matched_pdf["cert_no"]
            cert_cell = ws.cell(row=r, column=cert_col_idx)
            cert_cell.value = cert_val

            # Enforce Black font color ("000000") and Non-Bold font weight (bold=False)
            curr_font = cert_cell.font
            if curr_font:
                cert_cell.font = Font(
                    name=curr_font.name or "Arial",
                    size=curr_font.size or 10,
                    bold=False,
                    italic=curr_font.italic,
                    color="000000"
                )
            else:
                cert_cell.font = Font(color="000000", bold=False)

            rows_updated += 1
            matched_pdf["matched_excel"] = True
            
            # Determine sequence prefix (e.g. 001, 002) from Col 1 "No."
            seq_val = ws.cell(row=r, column=no_col_idx).value
            try:
                if seq_val is not None and str(seq_val).strip() != "":
                    seq_num = int(float(str(seq_val).strip()))
                    seq_prefix = f"{seq_num:03d}"
                else:
                    seq_prefix = f"{rows_updated:03d}"
            except Exception:
                seq_prefix = f"{rows_updated:03d}"

            new_pdf_filename = f"{seq_prefix}_{sanitize_filename(cert_val)}.pdf"
            matched_pdf["new_filename"] = new_pdf_filename
            matched_pdf["item_no"] = seq_prefix

            mapping_records.append({
                "Row": r,
                "No.": seq_prefix,
                "Serial No": cell_serial_val,
                "Certificate No": cert_val,
                "Original PDF": matched_pdf["filename"],
                "New PDF Name": new_pdf_filename,
                "Status": "✅ Matched & Mapped"
            })
        else:
            # If Serial No exists in Excel but no matching PDF was found, apply Yellow Fill to the ENTIRE ROW
            if cell_serial_val:
                max_cols_to_fill = max(ws.max_column, 9)
                for c in range(1, max_cols_to_fill + 1):
                    ws.cell(row=r, column=c).fill = yellow_fill

            mapping_records.append({
                "Row": r,
                "No.": str(ws.cell(row=r, column=no_col_idx).value or "-"),
                "Serial No": cell_serial_val,
                "Certificate No": str(ws.cell(row=r, column=cert_col_idx).value or ""),
                "Original PDF": matched_pdf["filename"] if matched_pdf else "-",
                "New PDF Name": "-",
                "Status": "⚠️ Unmatched Serial No (Entire Row Highlighted Yellow)" if cell_serial_val else "Info Row"
            })

    # 4. Inject IWK Logo into Excel Header (20% reduced size = 57.6 pt)
    logo_path = os.path.join("My Picture", "logo IWK.png")
    if os.path.exists(logo_path):
        try:
            ws.HeaderFooter.oddHeader.left.text = "&G"
            ws.HeaderFooter.oddHeader.left.image = logo_path
            logger.info("Embedded Logo in Excel openpyxl Left Header")
        except Exception as e:
            logger.warning(f"Could not set openpyxl left header image: {e}")

    # Save updated Excel bytes
    updated_excel_io = io.BytesIO()
    wb.save(updated_excel_io)
    updated_excel_bytes = updated_excel_io.getvalue()

    # Apply Excel COM automation header image with 20% reduced width (57.6 pt)
    if os.path.exists(logo_path):
        try:
            from backend.calibration_service import apply_header_image_with_com
            updated_excel_bytes = apply_header_image_with_com(updated_excel_bytes, logo_path, width=57.6)
            logger.info("Successfully applied header logo (57.6 pt width) via Excel COM automation")
        except Exception as com_err:
            logger.warning(f"Excel COM automation logo injection skipped or failed: {com_err}")

    # 5. Build ZIP file package (Include ONLY matched PDFs)
    zip_io = io.BytesIO()
    excel_basename = os.path.splitext(excel_filename)[0]
    updated_excel_filename = f"{excel_basename}_Mapped.xlsx"

    with zipfile.ZipFile(zip_io, "w", zipfile.ZIP_DEFLATED) as zip_file:
        # Add updated Excel file
        zip_file.writestr(updated_excel_filename, updated_excel_bytes)

        # Add ONLY matched PDFs (Exclude unmatched PDFs from download)
        for pdf_info in pdf_results:
            if pdf_info.get("matched_excel") and pdf_info.get("new_filename"):
                zip_file.writestr(f"Certificates/{pdf_info['new_filename']}", pdf_info["bytes"])

    zip_bytes = zip_io.getvalue()

    summary_df = pd.DataFrame([
        {
            "No.": p.get("item_no", "-"),
            "Original PDF Filename": p["filename"],
            "Extracted Serial No.": p["serial_no"] or "Not Found",
            "Extracted Certificate No.": p["cert_no"] or "Not Found",
            "New PDF Name": p.get("new_filename", "-"),
            "Mapping Status": "✅ Matched & Included in ZIP" if p.get("matched_excel") else ("⚠️ Unmatched (Excluded from ZIP)" if p["serial_no"] else "❌ Extraction Failed")
        }
        for p in pdf_results
    ])

    return {
        "updated_excel_bytes": updated_excel_bytes,
        "updated_excel_filename": updated_excel_filename,
        "zip_bytes": zip_bytes,
        "zip_filename": f"{excel_basename}_Certificates_Package.zip",
        "rows_updated": rows_updated,
        "total_pdfs": len(pdf_files_list),
        "matched_pdfs": sum(1 for p in pdf_results if p.get("matched_excel")),
        "summary_df": summary_df,
        "mapping_records": mapping_records
    }
