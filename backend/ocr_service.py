import os
import io
import re
import logging
import difflib
from typing import List, Dict, Any, Optional
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Header
from pydantic import BaseModel
import openpyxl
from PIL import Image
import fitz  # PyMuPDF

# Set up logging
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger("OCRService")

app = FastAPI(
    title="OCR Adjustment Service API",
    description="FastAPI service to extract handwritten values from scanned PDFs using Azure Document Intelligence.",
    version="1.0.0"
)

# ---------------------------------------------------------------------------
# Data Transfer Objects (Pydantic Models)
# ---------------------------------------------------------------------------

class ExtractedField(BaseModel):
    FieldName: str
    Value: str
    Confidence: float

class OCRExtractionResponse(BaseModel):
    DocumentName: str
    Status: str
    Fields: List[ExtractedField]

# Default fallback search labels if no Excel template is provided
DEFAULT_SEARCH_LABELS = [
    {"de": "Auftrags-Nr.", "en": "Order-no."},
    {"de": "Kunde", "en": "Customer"},
    {"de": "Bau-Nr.", "en": "Serial-no."},
    {"de": "Case discharge", "en": "Case discharge"},
    {"de": "Case turning unit", "en": "Case turning unit"},
    {"de": "Inserting pusher", "en": "Inserting pusher"},
    {"de": "Stacking pusher", "en": "Stacking pusher"}
]

# ---------------------------------------------------------------------------
# Core OCR & Spatial Helper Methods
# ---------------------------------------------------------------------------

def pdf_to_images(pdf_bytes: bytes) -> List[Image.Image]:
    """
    Converts PDF bytes into a list of PIL Images using PyMuPDF.
    """
    images = []
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            pix = page.get_pixmap(dpi=300)
            img_data = pix.tobytes("png")
            images.append(Image.open(io.BytesIO(img_data)))
        logger.info(f"Successfully converted PDF to {len(images)} page images.")
    except Exception as e:
        logger.error(f"Error converting PDF to images: {e}")
        raise HTTPException(status_code=500, detail=f"PDF rendering failed: {str(e)}")
    return images

def load_template_labels(excel_bytes: bytes) -> List[Dict[str, str]]:
    """
    Parses the "Master" sheet of the provided Excel template bytes to pull labels.
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), read_only=True)
        sheet = wb["Master"] if "Master" in wb.sheetnames else wb.active
        
        labels = []
        for r in range(1, min(sheet.max_row, 150) + 1):
            de_val = sheet.cell(row=r, column=2).value
            en_val = sheet.cell(row=r, column=4).value
            de_str = str(de_val).strip() if de_val is not None else ""
            en_str = str(en_val).strip() if en_val is not None else ""
            if de_str or en_str:
                labels.append({"de": de_str, "en": en_str})
        logger.info(f"Successfully loaded {len(labels)} label keys from uploaded Excel template.")
        return labels
    except Exception as e:
        logger.warning(f"Failed to parse Excel template labels: {e}. Falling back to default labels.")
        return DEFAULT_SEARCH_LABELS

def analyze_layout_with_azure(pdf_bytes: bytes, endpoint: str, key: str, page_image: Image.Image, page_idx: int) -> List[Dict[str, Any]]:
    """
    Calls Azure Document Intelligence layout API and maps coordinates to PIL image scale.
    """
    try:
        from azure.ai.formrecognizer import DocumentAnalysisClient
        from azure.core.credentials import AzureKeyCredential
    except ImportError:
        logger.error("Azure SDK is not installed.")
        raise HTTPException(status_code=500, detail="azure-ai-formrecognizer is not installed on server.")

    try:
        client = DocumentAnalysisClient(endpoint=endpoint, credential=AzureKeyCredential(key))
        poller = client.begin_analyze_document(
            model_id="prebuilt-layout",
            document=pdf_bytes,
            pages=f"{page_idx + 1}"
        )
        result = poller.result()
        
        if not result.pages:
            return []
            
        az_page = result.pages[0]
        img_w, img_h = page_image.size
        az_w, az_h = az_page.width, az_page.height
        
        scale_x = img_w / float(az_w)
        scale_y = img_h / float(az_h)
        
        blocks = []
        for line in az_page.lines:
            poly = line.polygon
            x_coords = []
            y_coords = []
            
            for pt in poly:
                if hasattr(pt, 'x'):
                    x_coords.append(pt.x)
                    y_coords.append(pt.y)
                elif isinstance(pt, (int, float)):
                    pass
                elif len(pt) == 2:
                    x_coords.append(pt[0])
                    y_coords.append(pt[1])
                    
            if not x_coords and isinstance(poly, list) and len(poly) >= 8:
                x_coords = [poly[i] for i in range(0, len(poly), 2)]
                y_coords = [poly[i+1] for i in range(0, len(poly), 2)]
                
            if x_coords and y_coords:
                x1 = min(x_coords) * scale_x
                y1 = min(y_coords) * scale_y
                x2 = max(x_coords) * scale_x
                y2 = max(y_coords) * scale_y
                
                blocks.append({
                    'bbox': [int(x1), int(y1), int(x2), int(y2)],
                    'text': line.content or getattr(line, 'text', ''),
                    'confidence': getattr(line, 'confidence', 0.88) or 0.88
                })
        return blocks
    except Exception as e:
        logger.error(f"Azure layout extraction failed: {e}")
        raise HTTPException(status_code=502, detail=f"Azure Document Intelligence error: {str(e)}")

def extract_fields_from_layout(blocks: List[Dict[str, Any]], labels: List[Dict[str, str]], search_limit=800) -> List[Dict[str, Any]]:
    """
    Spatially maps layout blocks to label anchors and extracts values to their right.
    """
    used_indices = set()
    extracted_fields = []
    
    for label_info in labels:
        de_lbl = label_info.get("de", "")
        en_lbl = label_info.get("en", "")
        lbl_display = de_lbl if de_lbl else en_lbl
        
        clean_de = re.sub(r'[^a-zA-Z0-9]', '', de_lbl).lower()
        clean_en = re.sub(r'[^a-zA-Z0-9]', '', en_lbl).lower()
        
        # Find best matching block anchor
        best_block_idx = None
        best_ratio = 0.0
        
        for idx, b in enumerate(blocks):
            if idx in used_indices:
                continue
            b_clean = re.sub(r'[^a-zA-Z0-9]', '', b['text']).lower()
            if len(b_clean) < 2:
                continue
                
            ratio_de = difflib.SequenceMatcher(None, b_clean, clean_de).ratio() if clean_de else 0.0
            ratio_en = difflib.SequenceMatcher(None, b_clean, clean_en).ratio() if clean_en else 0.0
            max_ratio = max(ratio_de, ratio_en)
            
            if max_ratio > best_ratio and max_ratio >= 0.75:
                best_ratio = max_ratio
                best_block_idx = idx
                
        if best_block_idx is not None:
            anchor = blocks[best_block_idx]
            ax1, ay1, ax2, ay2 = anchor['bbox']
            aw, ah = ax2 - ax1, ay2 - ay1
            
            # Search right for candidate values
            candidates = []
            for v_idx, v_b in enumerate(blocks):
                if v_idx == best_block_idx or v_idx in used_indices:
                    continue
                    
                vx1, vy1, vx2, vy2 = v_b['bbox']
                vw, vh = vx2 - vx1, vy2 - vy1
                
                # Check horizontal alignment
                is_aligned = (min(ay1 + ah, vy1 + vh) - max(ay1, vy1)) > 0 or abs((vy1 + vh/2) - (ay1 + ah/2)) < (ah * 1.5)
                is_right = vx1 >= (ax2 - 10)
                dist = vx1 - ax2
                
                if is_aligned and is_right and dist < search_limit:
                    candidates.append((dist, v_b, v_idx))
                    
            value_str = ""
            conf = anchor['confidence']
            if candidates:
                candidates.sort(key=lambda x: x[0])
                best_val_block = candidates[0][1]
                best_val_idx = candidates[0][2]
                
                value_str = best_val_block['text']
                conf = min(conf, best_val_block['confidence'])
                used_indices.add(best_val_idx)
                
            used_indices.add(best_block_idx)
            
            extracted_fields.append({
                "FieldName": lbl_display,
                "Value": value_str,
                "Confidence": conf
            })
        else:
            # Field not found on page
            extracted_fields.append({
                "FieldName": lbl_display,
                "Value": "",
                "Confidence": 0.0
            })
            
    return extracted_fields

# ---------------------------------------------------------------------------
# FastAPI API Endpoints
# ---------------------------------------------------------------------------

@app.post("/api/ocr/extract", response_model=OCRExtractionResponse)
async def extract_ocr_data(
    file: UploadFile = File(...),
    excel_template: Optional[UploadFile] = File(None),
    azure_endpoint: Optional[str] = Form(None),
    azure_key: Optional[str] = Form(None),
    x_azure_endpoint: Optional[str] = Header(None, alias="X-Azure-Endpoint"),
    x_azure_key: Optional[str] = Header(None, alias="X-Azure-Key")
):
    """
    Extracts text layout fields from a scanned PDF document.
    Allows dynamic templates via an uploaded Excel file.
    """
    logger.info(f"Received OCR extraction request for file: {file.filename}")
    
    # Resolve Azure credentials (Form parameters take precedence over headers)
    endpoint = azure_endpoint or x_azure_endpoint or os.getenv("AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT", "")
    key = azure_key or x_azure_key or os.getenv("AZURE_DOCUMENT_INTELLIGENCE_KEY", "")
    
    if not endpoint or not key:
        logger.error("Missing Azure Document Intelligence credentials.")
        raise HTTPException(
            status_code=400,
            detail="Azure Document Intelligence Endpoint and Key are required. Provide them in headers (X-Azure-Endpoint/X-Azure-Key) or form inputs."
        )
        
    try:
        pdf_bytes = await file.read()
        
        # Load Excel template labels if provided
        labels = DEFAULT_SEARCH_LABELS
        if excel_template:
            excel_bytes = await excel_template.read()
            labels = load_template_labels(excel_bytes)
            
        # Convert PDF to page images
        pages = pdf_to_images(pdf_bytes)
        
        all_extracted_fields = []
        
        # Process first page layout (most forms are single page templates)
        if len(pages) > 0:
            blocks = analyze_layout_with_azure(pdf_bytes, endpoint, key, pages[0], page_idx=0)
            fields = extract_fields_from_layout(blocks, labels)
            
            for f in fields:
                all_extracted_fields.append(
                    ExtractedField(
                        FieldName=f['FieldName'],
                        Value=f['Value'],
                        Confidence=f['Confidence']
                    )
                )
                
        return OCRExtractionResponse(
            DocumentName=file.filename,
            Status="Success",
            Fields=all_extracted_fields
        )
        
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Internal processing failure: {e}")
        raise HTTPException(status_code=500, detail=f"Internal server error: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("ocr_service:app", host="127.0.0.1", port=8000, reload=True)
